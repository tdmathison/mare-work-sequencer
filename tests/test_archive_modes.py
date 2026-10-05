import io,json,zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from werkzeug.security import generate_password_hash
from app import app,db,run,ensure_package,package
from reporting import TOKENS,fill_template

def test_archive_modes_previews_and_alignment(tmp_path):
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('archive-user',generate_password_hash('password-long-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('ARCH-1','20261005: Archive modes','2026',2)").lastrowid
        root=ensure_package(cid)
        (root/'reports/sections/assets').mkdir(parents=True)
        (root/'reports/sections/executive_summary.md').write_text('Editor text')
        (root/'reports/sections/assets/report-image-test.png').write_bytes(b'image')
        (root/'reports/malware-analysis-report.md').write_text('Working report')
        (root/'reports/backups').mkdir();(root/'reports/backups/old.docx').write_bytes(b'backup')
        doc=Document();doc.add_paragraph('Final report');doc.save(root/'reports/malware-analysis-report.docx')
    client=app.test_client()
    with client.session_transaction() as s:s.update(uid=uid,version=1,csrf='archive-token')
    def post(endpoint,**data):return client.post(f'/cases/{cid}'+endpoint,data={'csrf':'archive-token',**data})
    for kind in ('standard','raw'):
        response=post('/archive',archive_type=kind,acknowledge='1');assert response.status_code==200
        expected_suffix='-MIP-Archive_modes'+('-RAW' if kind=='raw' else '')+'.zip'
        assert response.headers['Content-Disposition'].endswith(expected_suffix)
        z=zipfile.ZipFile(io.BytesIO(response.data));names=z.namelist()
        assert any(n.endswith('malware-analysis-report.docx') for n in names)
        assert not any('/backups/' in n for n in names)
        assert any(n.endswith('reports/malware-analysis-report.md') for n in names)==(kind=='raw')
        assert any('/reports/sections/assets/' in n for n in names)==(kind=='raw')
        manifest=json.loads(z.read(next(n for n in names if n.endswith('manifest.json'))))
        assert manifest['archive_type']==kind
    page=client.get(f'/cases/{cid}').data.decode()
    standard=page.split('data-archive-pane="standard"')[1].split('data-archive-pane="raw"')[0]
    assert 'reports/malware-analysis-report.md' not in standard and 'reports/sections/assets' not in standard
    assert 'Download Standard MIP Archive' in page and 'Download RAW MIP Archive' in page
    assert 'class="word-wrap" checked' in page and 'markdown-block-style' in page and 'data-md="align-center"' in page
    package_view=page.split('id="panel-package"')[1].split('id="panel-indicators"')[0]
    assert 'reports/malware-analysis-report.md' not in package_view
    markdown='::: align-center\nCentered paragraph\n:::'
    html=post('/markdown-preview',text=markdown).json['html']
    assert 'align-center' in html and 'Centered paragraph' in html
    template=Document()
    for key,label in TOKENS.items():template.add_heading(label,1);template.add_paragraph('{{ '+key+' }}')
    path=tmp_path/'template.docx';template.save(path)
    destination=tmp_path/'out.docx'
    fill_template(path,destination,{key:markdown if key=='executive_summary' else 'Content' for key in TOKENS},'ARCH-1','Archive modes')
    paragraph=next(p for p in Document(destination).paragraphs if p.text=='Centered paragraph')
    assert paragraph.alignment==WD_ALIGN_PARAGRAPH.CENTER
