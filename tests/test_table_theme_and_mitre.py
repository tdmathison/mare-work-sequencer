import os,subprocess,sys,json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from reporting import TOKENS,fill_template
from mip import create_mip,OPTIONAL_TEMPLATES
from app import app,run,ensure_package
from werkzeug.security import generate_password_hash


def test_generated_tables_theme_and_template_preservation(tmp_path):
    template=tmp_path/'template.docx';output=tmp_path/'output.docx'
    d=Document();d.sections[0].header.paragraphs[0].text='Template branding'
    old=d.add_table(rows=1,cols=1);old.cell(0,0).text='Keep existing table';old.style='Table Grid'
    for key,label in TOKENS.items():d.add_heading(label,1);d.add_paragraph('{{ '+key+' }}')
    d.save(template)
    sections={key:'Analyst text' for key in TOKENS};sections['key_findings']='| Finding | Evidence |\n|---|---|\n| **Observed** | `Run key` |\n| Second | Evidence |'
    fill_template(template,output,sections,'MARE-1','Case',indicators={'columns':['type','value','description'],'rows':[['domain','example[.]com','C2']]},references={'rows':[['https://example.org','Reference']]})
    doc=Document(output)
    assert doc.sections[0].header.paragraphs[0].text=='Template branding'
    assert doc.tables[0].style.name=='Table Grid'
    assert len(doc.tables)==4
    for table in doc.tables[1:]:
        assert table.style.style_id=='MAREReportTable'
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        assert all(run.bold and run.font.color.rgb.__str__()=='FFFFFF' for cell in table.rows[0].cells for run in cell.paragraphs[0].runs)
    style=doc.styles['MARE Report Table'].element
    assert style.find('.//'+qn('w:shd')).get(qn('w:fill'))=='4F81BD'
    assert 'DBE5F1' in style.xml and '95B3D7' in style.xml and 'themeFill' not in style.xml


def test_new_and_existing_mapping_scaffolding(tmp_path):
    root=create_mip(tmp_path,'Case');assert list((root/'mappings').iterdir())==[]
    with app.app_context():
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('EMPTY-MAPPINGS','Empty','2026',1)").lastrowid
        root=ensure_package(cid)
        for rel,text in OPTIONAL_TEMPLATES.items():
            if rel.startswith('mappings/'):(root/rel).write_text(text)
        (root/'mappings/mitre-attack.md').write_text('### MITRE Attack\nAnalyst-reviewed mapping')
        ensure_package(cid)
        assert (root/'mappings/mitre-attack.md').read_text().endswith('Analyst-reviewed mapping')
        assert not (root/'mappings/mitre-mbc.md').exists()


def test_legacy_settings_migration_keeps_template(tmp_path):
    env={**os.environ,'MARE_DATA_DIR':str(tmp_path/'data')}
    subprocess.run([sys.executable,'-c',"from app import app,db,run,ROOT\nwith app.app_context():\n uid=run(\"INSERT INTO users(username,password,role) VALUES('legacy','hash','User')\").lastrowid\n run(\"CREATE TABLE ai_settings(user_id INTEGER PRIMARY KEY,key_cipher TEXT,model TEXT,template_id TEXT)\")\n run(\"INSERT INTO ai_settings VALUES(?,?,?,?)\",(uid,'encrypted-key','legacy-model','selected-template'))\n (ROOT/'api-keys.key').write_text('legacy key')"],env=env,check=True,capture_output=True)
    result=subprocess.run([sys.executable,'-c',"import json\nfrom app import app,db,ROOT\nwith app.app_context():\n print(json.dumps({'template':db().execute('SELECT template_id FROM report_settings').fetchone()[0],'ai_table':bool(db().execute(\"SELECT 1 FROM sqlite_master WHERE name='ai_settings'\").fetchone()),'key_exists':(ROOT/'api-keys.key').exists()}))"],env=env,check=True,capture_output=True,text=True)
    assert json.loads(result.stdout)=={'template':'selected-template','ai_table':False,'key_exists':False}
