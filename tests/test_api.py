import hashlib, io, json, zipfile
import pytest
from docx import Document
from werkzeug.security import generate_password_hash
from app import app, db, run, package, ensure_package
import report_routes
from reporting import TOKENS

PNG = None


def make_png():
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (4, 4), 'red').save(buffer, 'PNG')
    return buffer.getvalue()


def make_user(name, role='Administrator', forced=0):
    with app.app_context():
        uid = run('INSERT INTO users(username,password,role,forced) VALUES(?,?,?,?)', (name, generate_password_hash('long-password-123'), role, forced)).lastrowid
        secret = 'mare_test_' + name
        run('INSERT INTO api_tokens(user_id,name,token_hash,token_prefix,created) VALUES(?,?,?,?,?)', (uid, 'test', hashlib.sha256(secret.encode()).hexdigest(), secret[:14], '2026'))
    return uid, {'Authorization': 'Bearer ' + secret}


@pytest.fixture
def admin():
    app.config['TESTING'] = True
    return make_user('apiadmin')


def new_case(client, headers):
    response = client.post('/api/v1/cases', json={'name': 'API case'}, headers=headers)
    assert response.status_code == 201
    return response.json['case']['id']


def start_case(client, headers, cid, uid):
    assert client.patch(f'/api/v1/cases/{cid}', json={'owner_id': uid}, headers=headers).status_code == 200
    assert client.post(f'/api/v1/cases/{cid}/stage', json={'expected': 0}, headers=headers).status_code == 200


def word_template():
    document = Document()
    for key, label in TOKENS.items():
        document.add_heading(label, 1)
        document.add_paragraph('{{ ' + key + ' }}')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_authentication_and_errors(admin):
    client = app.test_client()
    assert client.get('/api/v1/me').status_code == 401
    assert client.get('/api/v1/me', headers={'Authorization': 'Bearer nope'}).status_code == 401
    missing = client.get('/api/v1/nope', headers=admin[1])
    assert missing.status_code == 404 and 'error' in missing.json
    assert client.get('/api/v1/me', headers=admin[1]).json['user']['username'] == 'apiadmin'
    forced_uid, forced_headers = make_user('forcedapi', forced=1)
    assert client.get('/api/v1/me', headers=forced_headers).status_code == 403


def test_user_role_permissions_and_expiry():
    app.config['TESTING'] = True
    with app.app_context():
        db().execute("INSERT INTO roles(name,description,builtin) VALUES('User','',1)")
        db().commit()
    uid, headers = make_user('plainapi', role='User')
    client = app.test_client()
    assert client.get('/api/v1/cases', headers=headers).status_code == 403
    admin_uid, admin_headers = make_user('expiryadmin')
    created = client.post('/api/v1/tokens', json={'name': 'short', 'expires_in_days': 1}, headers=admin_headers)
    assert created.status_code == 201 and created.json['expires']
    assert client.post('/api/v1/tokens', json={'name': 'bad', 'expires_in_days': 0}, headers=admin_headers).status_code == 400
    token = created.json['token']
    with app.app_context():
        run("UPDATE api_tokens SET expires='2000-01-01T00:00:00+00:00' WHERE id=?", (created.json['id'],))
    assert client.get('/api/v1/me', headers={'Authorization': 'Bearer ' + token}).status_code == 401


def test_rate_limit(admin, monkeypatch):
    monkeypatch.setenv('MARE_API_RATE_LIMIT', '2')
    import app as application
    application.API_HITS.clear()
    client = app.test_client()
    codes = [client.get('/api/v1/me', headers=admin[1]).status_code for _ in range(3)]
    assert codes == [200, 200, 429]
    application.API_HITS.clear()


def test_pagination_and_case_detail(admin):
    client = app.test_client()
    uid, headers = admin
    cid = new_case(client, headers)
    start_case(client, headers, cid, uid)
    for index in range(3):
        assert client.post(f'/api/v1/cases/{cid}/tasks', json={'title': f'Task {index}'}, headers=headers).status_code == 201
    listing = client.get(f'/api/v1/cases/{cid}/tasks?limit=2&offset=1', headers=headers).json
    assert listing['total'] == 3 and len(listing['tasks']) == 2
    assert client.get(f'/api/v1/cases/{cid}/tasks?limit=x', headers=headers).status_code == 400
    detail = client.get(f'/api/v1/cases/{cid}', headers=headers).json['case']
    assert detail['history'] and 'readiness' in detail and 'deferred' in detail and detail['owner_active']
    assert client.get('/api/v1/cases', headers=headers).json['total'] == 1
    assert client.get('/api/v1/users?limit=1', headers=headers).json['total'] >= 1


def test_templates_report_images_and_final_report(admin, monkeypatch):
    client = app.test_client()
    uid, headers = admin
    cid = new_case(client, headers)
    start_case(client, headers, cid, uid)
    assert client.post('/api/v1/templates', headers=headers, data={'template': (io.BytesIO(b'not a docx'), 'bad.docx')}, content_type='multipart/form-data').status_code == 400
    uploaded = client.post('/api/v1/templates', headers=headers, data={'template': (io.BytesIO(word_template()), 'mine.docx')}, content_type='multipart/form-data')
    assert uploaded.status_code == 201 and uploaded.json['templates'][0]['selected']
    tid = uploaded.json['templates'][0]['id']
    assert client.get('/api/v1/templates/' + tid, headers=headers).status_code == 200
    assert client.put('/api/v1/templates/selected', json={'template_id': 'missing'}, headers=headers).status_code == 404
    assert client.put('/api/v1/templates/selected', json={'template_id': ''}, headers=headers).json['templates'][0]['selected'] is False

    image = client.post(f'/api/v1/cases/{cid}/report/images', headers=headers, data={'section': 'executive_summary', 'image': (io.BytesIO(make_png()), 'shot.png')}, content_type='multipart/form-data')
    assert image.status_code == 201 and image.json['figure'] == 1
    name = image.json['name']
    assert client.get(image.json['url'], headers=headers).mimetype == 'image/png'
    assert [row['name'] for row in client.get(f'/api/v1/cases/{cid}/report/images', headers=headers).json['images']] == [name]
    assert client.delete(f'/api/v1/cases/{cid}/report/images/{name}', headers=headers).status_code == 204
    assert client.get(f'/api/v1/cases/{cid}/report/images/{name}', headers=headers).status_code == 404

    assert client.get(f'/api/v1/cases/{cid}/report/download', headers=headers).status_code == 404
    assert client.get(f'/api/v1/cases/{cid}/report/download?format=zip', headers=headers).status_code == 400
    document = Document()
    document.add_paragraph('Edited report')
    buffer = io.BytesIO()
    document.save(buffer)
    assert client.post(f'/api/v1/cases/{cid}/report/final', headers=headers, data={'report': (io.BytesIO(b'junk'), 'r.docx')}, content_type='multipart/form-data').status_code == 400
    final = client.post(f'/api/v1/cases/{cid}/report/final', headers=headers, data={'report': (io.BytesIO(buffer.getvalue()), 'r.docx')}, content_type='multipart/form-data')
    assert final.status_code == 200 and final.json['saved']
    assert client.get(f'/api/v1/cases/{cid}/report/download', headers=headers).data == buffer.getvalue()
    assert client.get(f'/api/v1/cases/{cid}/report/download?format=pdf', headers=headers).data.startswith(b'%PDF')

    assert client.delete('/api/v1/templates/' + tid, headers=headers).status_code == 204
    assert client.delete('/api/v1/templates/' + tid, headers=headers).status_code == 404


def test_backup_import_export_and_metrics(admin):
    client = app.test_client()
    uid, headers = admin
    cid = new_case(client, headers)
    start_case(client, headers, cid, uid)
    single = client.post(f'/api/v1/cases/{cid}/backup', headers=headers)
    assert single.status_code == 200 and zipfile.is_zipfile(io.BytesIO(single.data))
    everything = client.post('/api/v1/backups/export', headers=headers)
    assert everything.status_code == 200 and 'backup.json' in zipfile.ZipFile(io.BytesIO(everything.data)).namelist()
    assert client.post('/api/v1/backups/import', headers=headers).status_code == 400
    bad = client.post('/api/v1/backups/import', headers=headers, data={'backup': (io.BytesIO(b'garbage'), 'x.zip')}, content_type='multipart/form-data')
    assert bad.status_code == 400 and 'error' in bad.json
    restored = client.post('/api/v1/backups/import', headers=headers, data={'backup': (io.BytesIO(everything.data), 'all.zip')}, content_type='multipart/form-data')
    assert restored.status_code == 201, restored.json
    assert client.get('/api/v1/cases', headers=headers).json['total'] == 2
    raw = client.post('/api/v1/backups/import-raw', headers=headers, data={'backup': (io.BytesIO(single.data), 'one.zip')}, content_type='multipart/form-data')
    assert raw.status_code == 201, raw.json
    database = client.post('/api/v1/backups/database', headers=headers)
    assert database.status_code == 200 and 'data/workflow.sqlite3' in zipfile.ZipFile(io.BytesIO(database.data)).namelist()
    metrics = client.get('/api/v1/metrics', headers=headers).json
    assert {'users', 'stages', 'completed'} <= set(metrics) and metrics['users'][0]['username'] == 'apiadmin'


def test_api_page_lists_endpoints_and_tokens(admin):
    client = app.test_client()
    with client.session_transaction() as s:
        s.update(uid=admin[0], version=1, csrf='t')
    page = client.get('/api').data
    assert b'/cases/&lt;int:case_id&gt;/report/final' in page or b'/cases/<int:case_id>/report/final' in page
    assert b'import requests' in page and b'href="/api"' in page
    response = client.post('/account/api-tokens', data={'csrf': 't', 'name': 'ui', 'expires_in_days': '7'})
    assert response.status_code == 302 and response.headers['Location'].endswith('/api#api-tokens')
    assert b'Copy this token now' in client.get('/api').data
    assert b'api-tokens' not in client.get('/account').data


def test_case_name_gets_date_prefix(admin):
    from datetime import datetime
    client = app.test_client()
    today = datetime.now().strftime('%Y%m%d')
    plain = client.post('/api/v1/cases', json={'name': 'Plain'}, headers=admin[1]).json['case']['name']
    dated = client.post('/api/v1/cases', json={'name': '20200102: Old'}, headers=admin[1]).json['case']['name']
    assert plain == today + ': Plain' and dated == '20200102: Old'


def test_title_date_drives_report_and_package_names():
    from datetime import datetime, timezone
    from mip import title_date
    from reporting import report_filename
    today = datetime.now(timezone.utc).strftime('%Y%m%d')
    assert title_date('20200102: Old') == '20200102'
    assert title_date('No date') == today and title_date('20201345: Bad') == today
    case_row = {'name': '20200102: Old', 'number': 'MARE-2026-000007', 'external_number': ''}
    assert report_filename(case_row).startswith('20200102-MARE_000007_RE_Report_Old')
    assert report_filename({**case_row, 'name': 'Plain'}).startswith(today + '-')
