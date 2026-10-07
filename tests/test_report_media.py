import io,json,subprocess,zipfile
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches
from werkzeug.security import generate_password_hash
from app import app,db,run,ensure_package,package
from reporting import TOKENS,OUTPUT
import report_routes

def test_clipboard_image_paste_uses_editor_drop_upload():
    source=Path('static/clipboard-images.js').read_text()
    script='''const virtualMachine=require('vm'),assertions=require('assert');let pasteHandler,dropEvent;
    const area={closest:selector=>selector==='fieldset:disabled'?null:null,dispatchEvent:event=>dropEvent=event};
    const context={document:{addEventListener:(name,handler)=>{if(name==='paste')pasteHandler=handler;}},Event:function(type,options){this.type=type;Object.assign(this,options);}};
    virtualMachine.runInNewContext(SOURCE,context);
    let prevented=false;pasteHandler({target:{closest:()=>area},clipboardData:{types:['image/png'],items:[{kind:'file',type:'image/png',getAsFile:()=>({name:'screenshot.png'})}]},preventDefault:()=>prevented=true});
    assertions.equal(prevented,true);assertions.equal(dropEvent.type,'drop');assertions.equal(dropEvent.dataTransfer.files[0].name,'screenshot.png');
    prevented=false;pasteHandler({target:{closest:()=>area},clipboardData:{items:[{kind:'string',type:'text/plain'}]},preventDefault:()=>prevented=true});assertions.equal(prevented,false);
    prevented=false;pasteHandler({target:{closest:()=>area},clipboardData:{types:['text/plain','text/html','image/png'],items:[{kind:'string',type:'text/plain'},{kind:'file',type:'image/png',getAsFile:()=>({name:'word-clipboard-image.png'})}]},preventDefault:()=>prevented=true});assertions.equal(prevented,false);'''.replace('SOURCE',json.dumps(source))
    subprocess.run(['node','-e',script],check=True,capture_output=True,text=True)

def test_report_images_preview_word_zip_and_delete(monkeypatch,tmp_path):
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('media-user',generate_password_hash('password-long-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('MEDIA-1','Media','2026',1)").lastrowid
        ensure_package(cid)
    client=app.test_client()
    with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='media-token')
    base=f'/cases/{cid}'
    def post(endpoint,**data):return client.post(base+endpoint,data={'csrf':'media-token',**data})
    image=io.BytesIO();Image.new('RGB',(120,60),'red').save(image,format='PNG',dpi=(300,300))
    result=post('/report-images',section='executive_summary',image=(io.BytesIO(image.getvalue()),'screenshot.png'))
    assert result.status_code==200
    name=result.json['name'];relative=result.json['markdown_url'];path=package(cid)/'reports/sections'/relative
    assert name!='screenshot.png' and path.exists()
    with Image.open(path) as saved_image:assert saved_image.size==(120,60) and saved_image.info['dpi'][0]>290
    assert client.get(base+'/report-images/'+name).mimetype=='image/png'
    assert post('/report-images',section='executive_summary',image=(io.BytesIO(b'not image'),'bad.png')).status_code==400
    assert result.json['figure']==1
    text='Summary\n\n::: figure\n![Screenshot]('+relative+')\n\n*Figure 1:* \n:::\n\n|A|B|\n|---|---|\n|1|2|\n\n~~removed~~\n\n- [x] Done'
    preview=post('/markdown-preview',text=text).json['html']
    assert 'class="figure"' in preview and '<em>Figure 1:</em>' in preview
    assert '<table>' in preview and '<s>removed</s>' in preview and 'checkbox' in preview
    assert base+'/report-images/'+name in preview
    document=Document()
    for key,label in TOKENS.items():document.add_heading(label,1);document.add_paragraph('{{ '+key+' }}')
    template=tmp_path/'template.docx';document.save(template)
    assert client.post('/account/templates',data={'csrf':'media-token','template':(io.BytesIO(template.read_bytes()),'template.docx')}).status_code==302
    class Thread:
        def __init__(self,target,args,**kwargs):self.target=target;self.args=args
        def start(self):self.target(*self.args)
    monkeypatch.setattr(report_routes.threading,'Thread',Thread)
    assert post('/generate-report',executive_summary=text).status_code==202
    assert client.get(base+'/report-job').json['job']['status']=='completed'
    report=Document(package(cid)/client.get(f'/cases/{cid}/report-job').json['report_path'])
    assert len(report.inline_shapes)==1
    assert report.inline_shapes[0].width==Inches(120/300)
    assert '444444' in report.inline_shapes[0]._inline.xml
    caption=next(p for p in report.paragraphs if p.text=='Figure 1:')
    assert caption.alignment==1 and any(r.italic for r in caption.runs)
    assert not (package(cid)/'reports/indicators-of-compromise.md').exists()
    post('/stage',expected='1')
    archive=post('/archive',acknowledge='1');assert archive.status_code==200
    names=zipfile.ZipFile(io.BytesIO(archive.data)).namelist()
    assert not any(n.endswith('reports/sections/'+relative) for n in names)
    raw=post('/archive',archive_type='raw',acknowledge='1')
    raw_names=zipfile.ZipFile(io.BytesIO(raw.data)).namelist()
    assert any(n.endswith('reports/sections/'+relative) for n in raw_names)
    assert any(n.endswith('reports/sections/executive_summary.md') for n in raw_names)
    assert not any(n.endswith('executive_summary.md') or n.endswith('indicators-of-compromise.md') for n in names)
    assert post('/report-images/'+name+'/delete').status_code==200
    assert not path.exists() and name not in (package(cid)/'reports/sections/executive_summary.md').read_text()
    asset=post('/upload',category='scripts',ajax='1',file=(io.BytesIO(b'print(1)'),'decode.py'))
    assert asset.json['saved'] and asset.json['path']=='scripts/decode.py'
    assert client.get(base+'/assets?category=scripts').json['files'][0]['path']=='scripts/decode.py'
    assert post('/assets/delete',path='scripts/decode.py').json['deleted']
    assert client.get(base+'/assets?category=scripts').json['files']==[]
    assert post('/assets/delete',path='../../session.key').status_code==400
    page=client.get(base).data
    assert b'line-numbers' in page and b'data-md="table"' in page and b'Asset manager' in page

def test_mitre_mapping_editors_save_canonical_report_sources():
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('mitre-editor',generate_password_hash('password-long-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('MITRE-EDITOR','MITRE editor','2026',1)").lastrowid
        root=ensure_package(cid)
    client=app.test_client()
    with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='mitre-editor-token')
    assert client.get(f'/cases/{cid}/report-data').json['mitre_mappings']=={'attack':'','mbc':''}
    attack='### MITRE Attack\n\n| Technique | Evidence |\n|---|---|\n| T1059 | Script execution |'
    mbc='### MITRE MBC\n\n| Behavior | Evidence |\n|---|---|\n| B0001 | Observed |'
    response=client.post(f'/cases/{cid}/mitre-mappings',data={'csrf':'mitre-editor-token','mitre_attack_markdown':attack,'mitre_mbc_markdown':mbc})
    assert response.status_code==200 and response.json['saved']
    assert (root/'mappings/mitre-attack.md').read_text()==attack
    assert (root/'mappings/mitre-mbc.md').read_text()==mbc
    assert client.get(f'/cases/{cid}/report-data').json['mitre_mappings']=={'attack':attack,'mbc':mbc}

def test_mitre_mapping_save_has_autosave_countdown_and_manual_reset():
    source=Path('static/report-workspace.js').read_text()
    assert "mappingCountdown.id='save-mitre-mappings-countdown'" in source
    assert 'mitreAutosaveTimer=window.MareAutosave.register' in source
    assert 'canSave:()=>mitreMappingsDirty&&!mitreMappingsSaving&&!loading' in source
    assert 'mitreAutosaveTimer.reset();saveMitreMappings()' in source

def test_asset_manager_hides_and_protects_managed_report_markdown():
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('report-assets',generate_password_hash('password-long-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('REPORT-ASSETS','Report assets','2026',1)").lastrowid
        root=ensure_package(cid)
        (root/'reports/sections').mkdir(exist_ok=True)
        editor_paths=[f'reports/sections/{key}.md' for key in ('executive_summary','key_findings','detection_opportunities','reverse_engineering_findings')]
        for path in editor_paths:(root/path).write_text('Managed report editor content')
        (root/'reports/indicators-of-compromise.md').write_text('Generated indicators')
        (root/'reports/custom-analysis.md').write_text('# User-authored report notes')
        (root/'reports/sections/user-authored.md').write_text('# User-authored section notes')
    client=app.test_client()
    with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='report-assets-token')
    base=f'/cases/{cid}'
    def post(endpoint,**data):return client.post(base+endpoint,data={'csrf':'report-assets-token',**data})
    listed={item['path'] for item in client.get(base+'/assets?category=reports').json['files']}
    protected=editor_paths+['reports/indicators-of-compromise.md','reports/executive-summary.md','reports/malware-analysis-report.md']
    assert not listed.intersection(protected)
    assert {'reports/custom-analysis.md','reports/sections/user-authored.md'}<=listed
    assert post('/assets/delete',path='reports/sections/user-authored.md').json['deleted']
    for rel in protected:
        assert post('/assets/delete',path=rel).status_code==400
        assert (root/rel).is_file()
    (root/'reports/executive-summary.md').write_text('# Analyst executive summary')
    (root/'reports/malware-analysis-report.md').write_text('# Analyst final report')
    listed={item['path'] for item in client.get(base+'/assets?category=reports').json['files']}
    assert {'reports/executive-summary.md','reports/malware-analysis-report.md'}<=listed
