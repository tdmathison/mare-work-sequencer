import io,json,shutil,subprocess,zipfile
from pathlib import Path
from docx import Document
from werkzeug.security import generate_password_hash
from app import app,db,run,package,ROOT,ensure_package
from reporting import TOKENS,OUTPUT,template_sections,defang,convert_docx_to_pdf
import report_routes

def test_docx_to_pdf_uses_headless_libreoffice(monkeypatch,tmp_path):
    source=tmp_path/'report.docx';source.write_bytes(b'docx')
    destination=tmp_path/'report.pdf'
    monkeypatch.setattr(shutil,'which',lambda name:'/usr/bin/soffice' if name=='soffice' else None)
    def fake_run(command,**kwargs):
        output=Path(command[command.index('--outdir')+1])
        (output/(source.stem+'.pdf')).write_bytes(b'%PDF-1.4\n%%EOF\n')
        assert '--headless' in command and kwargs['timeout']==120
        return subprocess.CompletedProcess(command,0,'','')
    monkeypatch.setattr(subprocess,'run',fake_run)
    convert_docx_to_pdf(source,destination)
    assert destination.read_bytes().startswith(b'%PDF-')

def test_report_pipeline_and_private_settings(monkeypatch,tmp_path):
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('reporter',generate_password_hash('long-password-123'),'User')).lastrowid
        other=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('otherreporter',generate_password_hash('long-password-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage,owner) VALUES('REPORT-001','Reporting test','2026',1,?)",(uid,)).lastrowid
        root=ensure_package(cid)
    template=tmp_path/'template.docx';d=Document();d.add_heading('Analyst Report',0)
    for key,label in TOKENS.items():d.add_heading(label,1);d.add_paragraph('{{ '+key+' }}')
    d.save(template);assert template_sections(template)==set(TOKENS)
    client=app.test_client()
    with client.session_transaction() as s:s.update(uid=uid,version=1,csrf='test-token')
    def post(url,**data):return client.post(url,data={'csrf':'test-token',**data})
    post(f'/cases/{cid}/tasks',title='Review persistence',body='Check the evidence and record the offset.')
    with app.app_context():task=db().execute('SELECT * FROM tasks WHERE case_id=?',(cid,)).fetchone();assert task['body'].startswith('Check')
    post(f'/cases/{cid}/tasks/{task["id"]}/edit',title='Persistence review',body='Updated detail')
    assert post('/account/ai',model='test').status_code==404
    assert post('/account/templates',template=(io.BytesIO(template.read_bytes()),'My template.docx')).status_code==302
    with app.app_context():tid=db().execute('SELECT template_id FROM report_settings WHERE user_id=?',(uid,)).fetchone()[0]
    with client.session_transaction() as s:s['uid']=other
    assert client.get('/account/templates/'+tid).status_code==404
    with client.session_transaction() as s:s['uid']=uid
    (root/'supporting/sandbox.json').write_text('{"c2":"evil.example.com","ip":"192.0.2.5"}')
    backup=root/'reports/backups';backup.mkdir();(backup/'private-old.txt').write_text('PRIVATE BACKUP')
    sections={'executive_summary':'# Summary\nObserved persistence. Reference [research](https://example.org/paper).','key_findings':'Writes a Run key for persistence.','detection_opportunities':'Monitor relevant Run key changes.','reverse_engineering_findings':'**Configuration** uses `evil.example.com`.\n\n```python\nprint("decoded")\n```'}
    assert post(f'/cases/{cid}/report-sections',**sections).status_code==302
    data=client.get(f'/cases/{cid}/report-data').json
    assert all(data['sections'][k]==v for k,v in sections.items()) and all('backups' not in s['path'] for s in data['sources'])
    assert post(f'/cases/{cid}/indicators',columns=json.dumps(['type','value','description']),rows=json.dumps([['domain','evil.example.com','C2'],['ipv4','192.0.2.5','C2']]),revision='0').status_code==200
    fake={'attack':[{'id':'T1547.001','name':'Registry Run Keys / Startup Folder','evidence':'Writes a Run key for persistence.','source':'key_findings'}],'mbc':[],'iocs':[{'type':'domain','value':'evil.example.com','description':'Observed C2 endpoint','source':'supporting/sandbox.json'},{'type':'ipv4','value':'192.0.2.5','description':'Observed C2 address','source':'supporting/sandbox.json'}],'warnings':[]}
    (root/'mappings/mitre-attack.md').write_text('### MITRE Attack\n\n| ID | Technique |\n|---|---|\n| T1547.001 | Registry Run Keys |')
    class InlineThread:
        def __init__(self,target,args,**kwargs):self.target=target;self.args=args
        def start(self):self.target(*self.args)
    monkeypatch.setattr(report_routes.threading,'Thread',InlineThread)
    assert post(f'/cases/{cid}/generate-report',generation_mode='ai',**sections).status_code==400
    response=post(f'/cases/{cid}/generate-report',generation_mode='manual',**sections);assert response.status_code==202
    status=client.get(f'/cases/{cid}/report-job').json;assert status['job']['status']=='completed',status
    report_rel=client.get(f'/cases/{cid}/report-job').json['report_path'];report=root/report_rel;assert report.exists()
    contents=Document(report);text='\n'.join(p.text for p in contents.paragraphs)+'\n'+'\n'.join(c.text for t in contents.tables for row in t.rows for c in row.cells)
    assert 'evil[.]example[.]com' in text and '192[.]0[.]2[.]5' in text and 'T1547.001' in text
    assert '{{' not in text
    assert client.get(f'/cases/{cid}/file?path=reports/backups/private-old.txt').status_code==404
    assert b'private-old.txt' not in client.get(f'/cases/{cid}').data
    assert post(f'/cases/{cid}/generate-report',generation_mode='manual',**sections).status_code==409
    old=report.read_bytes()
    response=post(f'/cases/{cid}/generate-report',generation_mode='manual',previous='backup',**sections);assert response.status_code==202
    saved=list(backup.glob('*.docx'));assert len(saved)==1 and saved[0].read_bytes()==old
    post(f'/cases/{cid}/stage',expected='1')
    post(f'/cases/{cid}/stage',expected='2')
    archive=post(f'/cases/{cid}/archive',acknowledge='1');assert archive.status_code==200
    z=zipfile.ZipFile(io.BytesIO(archive.data));assert all('/backups/' not in n for n in z.namelist());assert any(n.endswith('.pdf') for n in z.namelist());assert not any(n.endswith('/'+report_rel) for n in z.namelist())
    # Preserve a Word file for internal render QA outside the deliverable archive.
    import shutil
    shutil.copy2(report,'/tmp/mare-report-qa.docx')

def test_evidence_validation_and_defanging():
    assert defang('https://evil.example.com/a')=='hxxps[:]//evil[.]example[.]com/a'
    assert defang('a'*64)=='a'*64

def test_manual_generation_and_final_word_upload(monkeypatch,tmp_path):
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('manualreporter',generate_password_hash('long-password-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage,owner) VALUES('MANUAL-001','Manual report','2026',1,?)",(uid,)).lastrowid
        root=ensure_package(cid)
    client=app.test_client()
    with client.session_transaction() as s:s.update(uid=uid,version=1,csrf='manual-token')
    def post(url,**data):return client.post(url,data={'csrf':'manual-token',**data})
    template=tmp_path/'manual-template.docx';d=Document();d.add_heading('Manual Report',0)
    for key,label in TOKENS.items():d.add_heading(label,1);d.add_paragraph('{{ '+key+' }}')
    d.save(template)
    post('/account/templates',template=(io.BytesIO(template.read_bytes()),'Manual template.docx'))
    class InlineThread:
        def __init__(self,target,args,**kwargs):self.target=target;self.args=args
        def start(self):self.target(*self.args)
    monkeypatch.setattr(report_routes.threading,'Thread',InlineThread)
    sections={key:'Analyst content for '+label for key,label in TOKENS.items()}
    sections['indicators_of_compromise']='| Type | Value | Description |\n|---|---|---|\n| Domain | manual[.]example | Analyst observed |'
    assert post(f'/cases/{cid}/generate-report',generation_mode='manual',**sections).status_code==202
    status=client.get(f'/cases/{cid}/report-job').json;assert status['job']['status']=='completed',status
    assert any('Formatting saved Markdown locally' in m for m in status['job']['messages'])
    with app.app_context():assert not db().execute("SELECT 1 FROM sqlite_master WHERE name='ai_settings'").fetchone()
    report_rel=client.get(f'/cases/{cid}/report-job').json['report_path'];report=root/report_rel;old=report.read_bytes();document=Document(report);document.add_paragraph('FINAL ANALYST REVISION');final=tmp_path/'edited.docx';document.save(final)
    assert report.with_suffix('.pdf').is_file()
    response=post(f'/cases/{cid}/final-report',report=(io.BytesIO(final.read_bytes()),'Edited report.docx'),backup='1');assert response.status_code==200
    assert report.read_bytes()==final.read_bytes() and report.with_suffix('.pdf').is_file()
    backups=list((root/'reports/backups').glob('*.docx'));assert len(backups)==1 and backups[0].read_bytes()==old
    page=client.get(f'/cases/{cid}').data;assert b'Open latest Word report' in page and b'Expand editor' in page
    download=client.get(f'/cases/{cid}/file?path={report_rel}&download=1');assert download.data==final.read_bytes()
    post(f'/cases/{cid}/stage',expected='1')
    post(f'/cases/{cid}/stage',expected='2')
    archive=post(f'/cases/{cid}/archive',acknowledge='1');assert archive.status_code==200
    z=zipfile.ZipFile(io.BytesIO(archive.data));assert z.read(next(n for n in z.namelist() if n.endswith('/'+report_rel)))==final.read_bytes();assert any(n.endswith('.pdf') for n in z.namelist()) and all('/backups/' not in n for n in z.namelist())
    import shutil
    shutil.copy2(report,'/tmp/mare-manual-report-qa.docx')
