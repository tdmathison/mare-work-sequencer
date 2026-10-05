import os, tempfile, io, zipfile, json, hashlib
os.environ['MARE_DATA_DIR']=tempfile.mkdtemp()
from app import app, db, run, package, ensure_package
from werkzeug.security import generate_password_hash

def test_workbench_lifecycle_and_security():
    app.config['TESTING']=True
    with app.app_context(): run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('admin',generate_password_hash('long-password-123'),'Administrator'))
    client=app.test_client()
    assert client.get('/').status_code==302
    assert client.post('/login',data={}).status_code==403
    client.get('/login')
    with client.session_transaction() as s: csrf=s['csrf']
    assert client.post('/login',data={'csrf':csrf,'username':'admin','password':'long-password-123'}).status_code==302
    with client.session_transaction() as s: csrf=s['csrf']
    def post(url,**data): return client.post(url,data={'csrf':csrf,**data})
    assert post('/cases',number='MARE-001',name='Test Case',description='<script>alert(1)</script>').status_code==302
    assert b'action="/cases"' not in client.get('/').data
    with app.app_context():
        first=db().execute('SELECT * FROM cases WHERE id=1').fetchone(); system_number=first['number']; assert system_number.startswith('MARE-') and system_number!='MARE-001'
    post('/cases/1',name='Test Case',number='tampered',description='<script>alert(1)</script>',external_system='JIRA',external_number='SEC-123',owner_id='1')
    with app.app_context():
        edited=db().execute('SELECT * FROM cases WHERE id=1').fetchone(); assert edited['number']==system_number and edited['external_number']=='SEC-123'
    assert post('/cases',name='Invalid reference',external_system='JIRA').status_code==400
    assert post('/cases',name='Second case',external_system='XSIAM',external_number='456').status_code==302
    with app.app_context():
        numbers=[r[0] for r in db().execute('SELECT number FROM cases')]; assert len(numbers)==len(set(numbers))
    assert (package(1)/'iocs').is_dir()
    assert not (package(1)/'samples').exists()
    for url in ['/','/cases/new','/cases/1','/users','/audit','/account']: assert client.get(url).status_code==200
    assert b'&lt;script&gt;' in client.get('/cases/1').data
    assert client.get('/cases/1/file?path=../../session.key').status_code==400
    assert post('/cases/1/draft',title='Locked',body='No edit').status_code==409
    assert post('/cases/1/stage',expected='0').status_code==302
    assert post('/cases/1/draft',title='Decode',body='Private draft').json['saved']
    assert b'Private draft' in client.get('/cases/1').data
    post('/cases/1/notes',title='Internal',body='INTERNAL SECRET')
    post('/cases/1/notes',title='Finding',body='EXPORTED FINDING',include_export='1')
    with app.app_context(): assert db().execute('SELECT count(*) FROM drafts').fetchone()[0]==0
    assert post('/cases/1/archive').status_code==409 # packaging is locked during analysis
    post('/cases/1/file',path='reports/malware-analysis-report.md',content='# Analysis\nObserved test findings.')
    from docx import Document
    word=Document();word.add_paragraph('Observed test findings.');word.save(package(1)/'reports/malware-analysis-report.docx')
    assert post('/cases/1/archive').status_code==409 # packaging is locked during analysis
    response=post('/cases/1/upload',category='scripts',file=(io.BytesIO(b'print(42)'),'decode.py'),description='Decode values',usage='python decode.py')
    assert response.status_code==302
    assert (package(1)/'scripts/decode.py').read_text()=='print(42)'
    post('/cases/1/upload',category='scripts',file=(io.BytesIO(b'overwrite'),'decode.py'))
    assert (package(1)/'scripts/decode.py').read_text()=='print(42)'
    assert post('/cases/1/upload',category='../',file=(io.BytesIO(b'x'),'x')).status_code==400
    for category in ['reports','scripts']: post('/cases/1/readiness',category=category,status='Populated')
    for category in ['iocs','signatures','mappings','supporting']: post('/cases/1/readiness',category=category,status='Not applicable',reason='No applicable findings')
    (package(1)/'samples').mkdir();(package(1)/'samples/payload.bin').write_bytes(b'payload')
    sections=package(1)/'reports/sections';sections.mkdir()
    (sections/'executive_summary.md').write_text('Editable report narrative')
    assert post('/cases/1/archive',acknowledge='1').status_code==409
    post('/cases/1/stage',expected='1',outcome='Deferred',reason='Awaiting source context')
    export=post('/cases/1/archive',acknowledge='1');assert export.status_code==200
    z=zipfile.ZipFile(io.BytesIO(export.data));names=z.namelist()
    assert not any('payload.bin' in n for n in names)
    assert not any('/reports/sections/' in n for n in names)
    assert (sections/'executive_summary.md').read_text()=='Editable report narrative'
    assert not any(n.endswith('reverse-engineering-notes.md') for n in names)
    raw=zipfile.ZipFile(io.BytesIO(post('/cases/1/archive',archive_type='raw',acknowledge='1').data))
    note=raw.read(next(n for n in raw.namelist() if n.endswith('reverse-engineering-notes.md'))).decode()
    assert 'EXPORTED FINDING' in note and 'INTERNAL SECRET' not in note
    manifest=json.loads(z.read(next(n for n in names if n.endswith('manifest.json'))))
    assert manifest['external_reference']=={'system':'JIRA','number':'SEC-123'}
    for f in manifest['files']:
        data=z.read(manifest['package_id']+'/'+f['path']);assert hashlib.sha256(data).hexdigest()==f['sha256']
    assert b'Usage:' not in z.read(next(n for n in names if n.endswith('README.md')))
    post('/cases/1/file',path='scripts/decode.py',content='print(43)')
    with app.app_context(): assert db().execute("SELECT status FROM readiness WHERE case_id=1 AND category='scripts'").fetchone()[0]=='Pending'
    assert post('/cases/1/archive').status_code==302
    assert post('/cases/1/archive',acknowledge='1').status_code==200
    post('/cases/1/readiness',category='scripts',status='Populated')
    post('/cases/1/stage',expected='2')
    with app.app_context(): assert db().execute('SELECT stage FROM cases WHERE id=1').fetchone()[0]==2
    post('/cases/1/deferred/1')
    post('/cases/1/stage',expected='2')
    with app.app_context(): assert db().execute('SELECT stage FROM cases WHERE id=1').fetchone()[0]==3
    assert post('/cases/1/notes',body='Blocked edit').status_code==409
    assert post('/cases/1/archive',acknowledge='1').status_code==200
    assert post('/cases/1/reopen').status_code==302
    assert post('/cases/1/notes',body='Reopened edit').status_code==302
    assert post('/users',username='analyst',role='User').status_code==302
    with app.app_context():
        u=db().execute("SELECT * FROM users WHERE username='analyst'").fetchone();uid=u['id'];assert u['forced']==1
    post(f'/users/{uid}/toggle')
    with client.session_transaction() as s:s['uid']=uid;s['version']=1
    assert client.get('/').status_code==302

def test_legacy_assets_preserved():
    with app.app_context():
        cid=run("INSERT INTO cases(number,name,description,created) VALUES('OLD','Old','','2026')").lastrowid
        root=package(cid);(root/'indicators').mkdir(parents=True);(root/'indicators/ips.txt').write_text('1.2.3.4')
        (root/'decoders').mkdir();(root/'decoders/decode.py').write_text('print(1)')
        ensure_package(cid)
        assert (root/'iocs/ips.txt').read_text()=='1.2.3.4'
        assert (root/'scripts/decode.py').read_text()=='print(1)'
        assert (root/'indicators/ips.txt').exists()
        (root/'iocs/ips.txt').write_text('new');ensure_package(cid)
        assert (root/'iocs/ips.txt').read_text()=='new'

def test_completion_blockers_visible_and_actionable():
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('completion-reviewer',generate_password_hash('review-password-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage,owner) VALUES('COMPLETE-001','Completion check','2026',2,?)",(uid,)).lastrowid
        root=ensure_package(cid)
        run('INSERT INTO tasks(case_id,title,body,created) VALUES(?,?,?,?)',(cid,'Review findings','Check final findings','2026'))
        run('INSERT INTO deferred VALUES(?,?,?,0)',(cid,1,'Awaiting evidence'))
    client=app.test_client()
    with client.session_transaction() as s:s.update(uid=uid,version=1,csrf='review-token')
    response=client.post(f'/cases/{cid}/stage',data={'csrf':'review-token','expected':'2'})
    assert response.status_code==302 and response.location.endswith('?tab=review#completion-review')
    page=client.get(response.location).data
    for text in [b'Completion checklist',b'Reports needs review',b'IOCs: Pending',b'Open task: Review findings',b'Deferred work remains for Malware Analysis',b'data-case-target="readiness-reports"',b'data-case-target="stage-history"',b'Downloading a ZIP with acknowledged issues']:
        assert text in page,text
    with app.app_context():assert db().execute('SELECT stage FROM cases WHERE id=?',(cid,)).fetchone()[0]==2

def test_merge_analysis_stage_migration():
    from app import migrate_analysis_stages
    with app.app_context():
        connection=db()
        for old in range(8):
            cid=run('INSERT INTO cases(number,name,created,stage) VALUES(?,?,?,?)',(f'MIG-{old}','Migration','2026',old)).lastrowid
            run('INSERT INTO history(case_id,stage) VALUES(?,?)',(cid,old))
            if old==3:
                run('INSERT INTO deferred VALUES(?,?,?,?)',(cid,2,'Static follow-up',1))
                run('INSERT INTO deferred VALUES(?,?,?,?)',(cid,3,'Dynamic follow-up',0))
        migrate_analysis_stages(connection)
        expected=[0,1,2,2,3,4,5,6]
        assert [r[0] for r in connection.execute('SELECT stage FROM cases ORDER BY id')]==expected
        assert [r[0] for r in connection.execute('SELECT stage FROM history ORDER BY id')]==expected
        deferred=connection.execute('SELECT * FROM deferred').fetchone()
        assert deferred['stage']==2 and deferred['resolved']==0
        assert 'Static follow-up' in deferred['reason'] and 'Dynamic follow-up' in deferred['reason']
        migrate_analysis_stages(connection)
        assert [r[0] for r in connection.execute('SELECT stage FROM cases ORDER BY id')]==expected

def test_five_stage_migration():
    from app import migrate_five_stages
    with app.app_context():
        connection=db()
        for old in range(7):
            cid=run('INSERT INTO cases(number,name,created,stage) VALUES(?,?,?,?)',(f'FIVE-{old}','Migration','2026',old)).lastrowid
            run('INSERT INTO history(case_id,stage) VALUES(?,?)',(cid,old))
            if old==2:
                for stage in (1,2,3):run('INSERT INTO deferred VALUES(?,?,?,?)',(cid,stage,f'Follow-up {stage}',int(stage!=2)))
        migrate_five_stages(connection)
        expected=[0,1,1,1,2,3,4]
        assert [r[0] for r in connection.execute('SELECT stage FROM cases ORDER BY id')]==expected
        assert [r[0] for r in connection.execute('SELECT stage FROM history ORDER BY id')]==expected
        row=connection.execute('SELECT * FROM deferred').fetchone()
        assert row['stage']==1 and row['resolved']==0
        assert all(f'Follow-up {stage}' in row['reason'] for stage in (1,2,3))
        migrate_five_stages(connection)
        assert [r[0] for r in connection.execute('SELECT stage FROM cases ORDER BY id')]==expected

def test_stage_permissions_and_ui():
    import re
    def element(page,identifier,tag='button'):
        return re.search(r'<'+tag+r'[^>]*id="'+identifier+r'"[^>]*>',page).group(0)
    def fieldset(page,identifier):
        return re.search(r'<form[^>]*id="'+identifier+r'"[^>]*>\s*(<fieldset[^>]*>)',page).group(1)
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('stage-user',generate_password_hash('stage-password-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage,owner) VALUES('LOCKS','Stage locks','2026',0,?)",(uid,)).lastrowid
        ensure_package(cid)
    client=app.test_client()
    with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='stage-token')
    def post(endpoint,**data):return client.post(f'/cases/{cid}'+endpoint,data={'csrf':'stage-token',**data})
    actions=['/draft','/notes','/tasks','/upload','/readiness','/generate-report','/final-report','/report-sections','/file','/mip','/artifact']
    for endpoint in actions:assert post(endpoint).status_code==409,endpoint
    page=client.get(f'/cases/{cid}').data.decode()
    assert 'disabled' in element(page,'tab-report')
    assert 'disabled' in fieldset(page,'note-form')
    assert 'Start analysis' in page
    post('/stage',expected='0')
    for stage in (1,):
        assert post('/archive',acknowledge='1').status_code==409
        assert post('/notes',body='Working note').status_code==302
        page=client.get(f'/cases/{cid}').data.decode()
        assert 'disabled' not in element(page,'tab-report')
        assert 'disabled' not in fieldset(page,'report-form')
        assert 'disabled' in re.search(r'<form[^>]*action="/cases/'+str(cid)+r'/archive"[^>]*>\s*(<fieldset[^>]*>)',page).group(1)
        post('/stage',expected=str(stage))
    with app.app_context():
        (package(cid)/'reports/malware-analysis-report.md').write_text('# Finished report')
        from docx import Document
        doc=Document();doc.add_paragraph('Finished report');doc.save(package(cid)/'reports/malware-analysis-report.docx')
        for category in ('reports','iocs','signatures','scripts','mappings','supporting'):
            run('INSERT OR REPLACE INTO readiness VALUES(?,?,?,?)',(cid,category,'Populated' if category=='reports' else 'Not applicable','No applicable output'))
    assert post('/archive').status_code==200
    assert post('/stage',expected='2').status_code==302
    for endpoint in actions:assert post(endpoint).status_code==409,endpoint
    assert post('',name='Locked rename').status_code==409
    assert post('/archive').status_code==200
    assert b'Reopen case' in client.get(f'/cases/{cid}').data
    assert post('/reopen').status_code==302
    assert post('/notes',body='Follow-up note').status_code==302
    with app.app_context():assert db().execute('SELECT stage FROM cases WHERE id=?',(cid,)).fetchone()[0]==1
