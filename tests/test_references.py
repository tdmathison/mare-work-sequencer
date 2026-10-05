import io,json
from docx import Document
from app import app,db,run,package,ensure_package
from reporting import TOKENS,OUTPUT,template_sections
from werkzeug.security import generate_password_hash
import report_routes

def test_references_saved_then_rendered_in_manual_word(monkeypatch,tmp_path):
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('refs-user',generate_password_hash('password-long-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('REF-1','References','2026',1)").lastrowid
        root=ensure_package(cid)
        legacy=root/'reports/sections';legacy.mkdir();(legacy/'threat_overview.md').write_text('Preserved legacy findings')
    client=app.test_client()
    with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='ref-token')
    def post(path,**data):return client.post(path,data={'csrf':'ref-token',**data})
    base=f'/cases/{cid}'
    assert client.get(base+'/report-data').json['sections']['key_findings']=='Preserved legacy findings'
    rows=[['https://example.org/research','Research **literal** | details'],['https://example.org/sandbox','Sandbox report']]
    assert post(base+'/references',rows=json.dumps(rows),revision='0').status_code==200
    assert post(base+'/references',rows='[]',revision='0').status_code==409
    assert len(client.get(base+'/references').json['rows'])==2
    # Delete a row and save.
    assert post(base+'/references',rows=json.dumps(rows[:1]),revision='1').status_code==200
    assert 'https://example.org/research' in (root/'supporting/references.csv').read_text()
    template=Document()
    for key,label in TOKENS.items():
        template.add_heading('Threat Overview' if key=='key_findings' else label,1)
        template.add_paragraph('{{ '+('threat_overview' if key=='key_findings' else key)+' }}')
    path=tmp_path/'legacy-template.docx';template.save(path)
    assert template_sections(path)==set(TOKENS)
    post('/account/templates',template=(io.BytesIO(path.read_bytes()),'legacy.docx'))
    class InlineThread:
        def __init__(self,target,args,**kwargs):self.target=target;self.args=args
        def start(self):self.target(*self.args)
    monkeypatch.setattr(report_routes.threading,'Thread',InlineThread)
    def never_ai(*args):raise AssertionError('UI/manual generation must not call AI')
    result=post(base+'/generate-report',executive_summary='Summary',key_findings='Key finding text',detection_opportunities='Detections',reverse_engineering_findings='RE notes')
    assert result.status_code==202
    assert client.get(base+'/report-job').json['job']['status']=='completed'
    report=Document(root/client.get(f'/cases/{cid}/report-job').json['report_path'])
    text='\n'.join(p.text for p in report.paragraphs)
    assert 'Key Findings' in text and 'Key finding text' in text and 'Threat Overview' not in text
    assert text.index('Links')>text.index('Appendices')
    links=next(t for t in report.tables if [c.text for c in t.rows[0].cells]==['Link','Description'])
    assert len(links.rows)==2 and links.rows[1].cells[1].text==rows[0][1]
    assert any(rel.target_ref==rows[0][0] for rel in report.part.rels.values() if rel.is_external)
    page=client.get(base).data
    for removed in (b'Create AI draft',b'AI writing assistance',b'Use AI to populate Word template',b'Ingest indicators from report',b'name="appendices"',b'name="threat_overview"'):assert removed not in page
    assert b'Generate Word Report' in page and b'tab-references' in page and b'name="key_findings"' in page
    assert b'OpenAI' not in client.get('/account').data
    with app.app_context():run('UPDATE cases SET stage=3 WHERE id=?',(cid,))
    assert post(base+'/references',rows='[]',revision='2').status_code==409
