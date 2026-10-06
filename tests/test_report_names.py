import io
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from app import app,db,run,ensure_package,package
from reporting import TOKENS,OUTPUT,report_filename,fill_template
from werkzeug.security import generate_password_hash
import report_routes

@pytest.mark.parametrize('external,title,expected',[
 ('12345','20261005: sqldeveloper64.exe','20261005-MARE_12345_RE_Report_sqldeveloper64_exe.docx'),
 ('','sqldeveloper64.exe','20261005-MARE_000001_RE_Report_sqldeveloper64_exe.docx'),
 ('JIRA/123','20261005: foo . / \\ bar__ test','20261005-MARE_JIRA_123_RE_Report_foo_bar_test.docx'),
 ('','Title without prefix','20261005-MARE_000001_RE_Report_Title_without_prefix.docx'),
])
def test_report_filename(external,title,expected):
 assert report_filename({'external_number':external,'number':'MARE-2026-000001','name':title},'20261005')==expected

@pytest.mark.parametrize('external_system,external_number,expected_number',[
 ('','', 'MARE #2026-000001'),
 ('Vortex','1234','Vortex #1234'),
 ('XSIAM','1234','XSIAM #1234'),
 ('JIRA','1234','JIRA #1234'),
])
def test_case_identity_template_placeholders(tmp_path,external_system,external_number,expected_number):
 template=tmp_path/'identity.docx';output=tmp_path/'filled.docx';document=Document()
 for key,label in TOKENS.items():document.add_heading(label,1);document.add_paragraph('{{ '+key+' }}')
 title_paragraph=document.add_paragraph();title_run=title_paragraph.add_run('{{ case_title }}');title_run.bold=True;title_run.font.color.rgb=RGBColor(255,255,255)
 shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'000000');title_paragraph._p.get_or_add_pPr().append(shading)
 number_paragraph=document.add_paragraph();number_paragraph.add_run('{{ case_');number_paragraph.add_run('number }}')
 document.sections[0].header.paragraphs[0].text='{{ case_title }} | {{ case_name }}';document.save(template)
 fill_template(template,output,{key:'Content' for key in TOKENS},'MARE-2026-000001','20261006: Malware X',external_system=external_system,external_number=external_number)
 result=Document(output);paragraphs=[p.text for p in result.paragraphs]
 assert 'Malware X' in paragraphs and expected_number in paragraphs
 formatted_title=next(p for p in result.paragraphs if p.text=='Malware X')
 assert formatted_title.runs[0].bold and formatted_title.runs[0].font.color.rgb==RGBColor(255,255,255)
 assert formatted_title._p.pPr.find(qn('w:shd')).get(qn('w:fill'))=='000000'
 assert result.sections[0].header.paragraphs[0].text=='Malware X | 20261006: Malware X'


def test_named_report_links_replacement_upload_and_backup_import(monkeypatch,tmp_path):
 with app.app_context():
  uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('namedreport',generate_password_hash('password-long-123'),'User')).lastrowid
  cid=run("INSERT INTO cases(number,name,created,stage,external_system,external_number) VALUES('MARE-2026-000001','20261005: sqldeveloper64.exe','2026',1,'JIRA','12345')").lastrowid
  root=ensure_package(cid)
 client=app.test_client()
 with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='named')
 def post(path,**data):return client.post(path,data={'csrf':'named',**data})
 d=Document()
 for key,label in TOKENS.items():d.add_heading(label,1);d.add_paragraph('{{ '+key+' }}')
 template=tmp_path/'template.docx';d.save(template)
 post('/account/templates',template=(io.BytesIO(template.read_bytes()),'template.docx'))
 class InlineThread:
  def __init__(self,target,args,**kwargs):self.target=target;self.args=args
  def start(self):self.target(*self.args)
 monkeypatch.setattr(report_routes.threading,'Thread',InlineThread)
 monkeypatch.setattr(report_routes,'report_filename',lambda c:report_filename(c,'20261005'))
 base=f'/cases/{cid}'
 assert post(base+'/generate-report',executive_summary='Findings').status_code==202
 status=client.get(base+'/report-job').json
 assert status['job']['status']=='completed',status
 rel=status['report_path'];assert rel=='reports/20261005-MARE_12345_RE_Report_sqldeveloper64_exe.docx'
 assert (root/rel).exists() and not (root/OUTPUT).exists()
 assert rel.encode() in client.get(base).data
 assert client.get(base+'/file',query_string={'path':rel,'download':'1'}).data==(root/rel).read_bytes()
 original=(root/rel).read_bytes()
 assert post(base+'/generate-report',executive_summary='New').status_code==409
 # Changed case title produces a new filename and backs up the previous report.
 with app.app_context():run("UPDATE cases SET name='20261005: new title.exe' WHERE id=?",(cid,))
 assert post(base+'/generate-report',executive_summary='New',previous='backup').status_code==202
 rel2=client.get(base+'/report-job').json['report_path']
 assert rel2.endswith('_new_title_exe.docx') and not (root/rel).exists()
 assert next((root/'reports/backups').glob('*.docx')).read_bytes()==original
 edited=Document(root/rel2);edited.add_paragraph('Final revision');final=tmp_path/'final.docx';edited.save(final)
 result=post(base+'/final-report',report=(io.BytesIO(final.read_bytes()),'Different download name.docx')).json
 assert result['report_path']==rel2 and (root/rel2).read_bytes()==final.read_bytes()
 # RAW backup retains the current report pointer despite allocation of a new case number.
 backup=post(base+'/backup');assert backup.status_code==200
 restored=post('/import-export/import-raw',backup=(io.BytesIO(backup.data),'named.zip'))
 assert restored.status_code==302
 with app.app_context():
  imported=db().execute('SELECT * FROM cases WHERE id!=? ORDER BY id DESC LIMIT 1',(cid,)).fetchone()
  assert imported['number']!='MARE-2026-000001' and imported['report_path']==rel2
  assert (package(imported['id'])/rel2).read_bytes()==final.read_bytes()
