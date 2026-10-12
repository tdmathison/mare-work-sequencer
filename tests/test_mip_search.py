import io
import json
import shutil
import stat
import zipfile
from pathlib import Path

import pytest
from docx import Document
from app import app, db, run, ROOT, ensure_package, build_archive, backup_helpers
from authorization import initialize_roles
from mip_search import extract, match_spans, MAX_BYTES


@pytest.fixture
def search_client(monkeypatch):
    index = app.extensions["mip_search"]
    # Keep assertions deterministic; production uses the background scheduler.
    monkeypatch.setattr(index, "schedule", lambda: None)
    shutil.rmtree(ROOT / "archives", ignore_errors=True)
    with index.connect() as con:
        con.execute("DELETE FROM documents")
    with app.app_context():
        initialize_roles(db())
        uid = run(
            "INSERT INTO users(username,password,role,forced) VALUES('analyst','unused','User',0)"
        ).lastrowid
        cid = run(
            "INSERT INTO cases(number,name,description,stage,created,owner,creator) VALUES('MARE-001','Search case','case needle',1,'now',?,?)",
            (uid, uid),
        ).lastrowid
        ensure_package(cid)
    client = app.test_client()
    with client.session_transaction() as session:
        session.update(uid=uid, version=1, csrf="search-token")
    return client, cid, index


def refresh(index):
    with app.app_context():
        index.refresh()
    assert index.status()["state"] == "ready", index.status()


def results(client, q, **filters):
    response = client.get("/search/results", query_string={"q": q, **filters})
    assert response.status_code == 200, response.data
    return response.json


def findings(result):
    return [f for group in result["groups"] for f in group["findings"]]


def test_database_fields_refang_and_filters(search_client):
    client, cid, index = search_client
    with app.app_context():
        run(
            "INSERT INTO notes(case_id,title,body,author) VALUES(?,'Named note','host evil[.]example[.]com','analyst')",
            (cid,),
        )
        run(
            "INSERT INTO tasks(case_id,title,body,priority,state) VALUES(?,'Find marker','task marker','High','In Progress')",
            (cid,),
        )
        sections = [
            {
                "id": "ioc-one",
                "title": "Indicator heading",
                "description": "section needle",
                "rows": [["domain", "evil[.]example[.]com", "row marker"]],
            }
        ]
        run("INSERT INTO indicator_sections VALUES(?,?)", (cid, json.dumps(sections)))
        run(
            "INSERT INTO indicator_tables VALUES(?,?,?,?)",
            (
                cid,
                json.dumps(["type", "value", "description"]),
                json.dumps(sections[0]["rows"]),
                1,
            ),
        )
        run(
            "INSERT INTO reference_tables VALUES(?,?,?)",
            (cid, json.dumps([["https://evil.example.com", "reference marker"]]), 1),
        )
    refresh(index)
    r = results(client, "evil.example.com")
    assert r["packages"] == 1 and r["artifacts"] == 3
    assert {f["kind"] for f in findings(r)} == {"notes", "indicators", "references"}
    assert results(client, "needle", kind="metadata")["artifacts"] == 1
    assert results(client, "marker", kind="tasks")["artifacts"] == 1
    assert results(client, "marker", stage="3")["artifacts"] == 0
    assert results(client, "marker", source="archive")["artifacts"] == 0
    assert client.get("/search/results?q=x&kind=invalid").status_code == 400
    assert client.get("/search/results?q=x&page=bad").status_code == 400
    assert results(client, "")["artifacts"] == 0


def test_text_filename_escaping_and_stale(search_client):
    client, cid, index = search_client
    path = ROOT / "cases" / str(cid) / "mip/supporting/hostile.txt"
    path.write_text('first needle\n<script>alert("needle")</script>\nneedle')
    refresh(index)
    r = results(client, "needle", kind="other")
    row = findings(r)[0]
    assert row["count"] == 3 and len(r["groups"]) == 1
    assert results(client, "hostile", kind="filenames")["artifacts"] == 1
    response = client.get("/search/findings/" + row["id"] + "?q=needle")
    assert response.status_code == 200
    assert b"<script>alert" not in response.data and b"&lt;script&gt;" in response.data
    assert b"/cases/" in response.data
    path.write_text("changed needle")
    assert client.get("/search/findings/" + row["id"]).status_code == 409
    refresh(index)
    assert client.get("/search/findings/" + row["id"]).status_code == 410
    path.unlink()
    refresh(index)
    assert results(client, "needle", kind="other")["artifacts"] == 0


def text_pdf(text):
    from pypdf import PdfWriter
    from pypdf.generic import (
        DictionaryObject,
        NameObject,
        NumberObject,
        DecodedStreamObject,
    )

    writer = PdfWriter()
    page = writer.add_blank_page(600, 800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_documents_and_extraction_cache(search_client, monkeypatch):
    client, cid, index = search_client
    folder = ROOT / "cases" / str(cid) / "mip/reports"
    doc = Document()
    doc.add_paragraph("paragraph marker")
    doc.add_table(1, 1).cell(0, 0).text = "table marker"
    doc.save(folder / "test.docx")
    (folder / "test.pdf").write_bytes(text_pdf("PDF marker"))
    refresh(index)
    assert results(client, "paragraph marker")["artifacts"] == 1
    assert results(client, "table marker")["artifacts"] == 1
    assert results(client, "PDF marker")["artifacts"] == 1
    import mip_search

    monkeypatch.setattr(
        mip_search,
        "extract",
        lambda *a: pytest.fail("Unchanged files must not be reparsed"),
    )
    refresh(index)
    assert results(client, "marker")["artifacts"] == 2


def archive(path, number, value, raw=False):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "pkg/manifest.json",
            json.dumps(
                {
                    "schema": "mare-mip/2.0",
                    "archive_type": "raw" if raw else "standard",
                    "case_number": number,
                    "title": "Archive case",
                }
            ),
        )
        z.writestr("pkg/supporting/evidence.txt", value)
        if raw:
            z.writestr(
                "pkg/case-data.json",
                json.dumps(
                    {
                        "schema": "mare-case-state/1",
                        "case": {"number": number},
                        "records": {
                            "notes": [
                                {
                                    "title": "Archived note",
                                    "body": "archive note needle",
                                }
                            ]
                        },
                    }
                ),
            )


def test_archive_versions_original_preview_and_unassociated(search_client):
    client, cid, index = search_client
    root = ROOT / "archives"
    root.mkdir()
    archive(root / "v1.zip", "MARE-001", "old version needle needle")
    archive(root / "v2.zip", "MARE-001", "new version needle", True)
    archive(root / "unassociated.zip", "OTHER", "unassociated needle")
    refresh(index)
    r = results(client, "needle", source="archive")
    assert r["packages"] == 1 and r["artifacts"] == 3
    old = next(f for f in findings(r) if f["version"].endswith("v1.zip"))
    assert old["count"] == 2
    (ROOT / "cases" / str(cid) / "mip/supporting/evidence.txt").write_text(
        "Current content changed"
    )
    response = client.get("/search/findings/" + old["id"])
    assert (
        b"old version needle needle" in response.data
        and b"Current content changed" not in response.data
    )
    with app.app_context():
        run("UPDATE users SET role='Administrator'")
    assert results(client, "needle", source="archive")["packages"] == 2
    archive(root / "v1.zip", "MARE-001", "replaced")
    assert client.get("/search/findings/" + old["id"]).status_code == 409


def test_real_standard_raw_archives_and_no_export_retention(search_client):
    client, cid, index = search_client
    with app.app_context():
        (ROOT / "cases" / str(cid) / "mip/supporting/proof.txt").write_text(
            "stored evidence needle"
        )
        run(
            "INSERT INTO notes(case_id,title,body,author,created,include_export) VALUES(?,'note','note archive needle','analyst','now',1)",
            (cid,),
        )
        root = ROOT / "archives"
        root.mkdir()
        for kind in ("standard", "raw"):
            out, name = (
                build_archive(cid, kind)
                if kind == "standard"
                else backup_helpers["case_backup"](cid)
            )
            with out:
                (root / (kind + ".zip")).write_bytes(out.read())
    refresh(index)
    assert results(client, "stored evidence needle", source="archive")["artifacts"] == 2
    assert results(client, "note archive needle", source="archive")["artifacts"] >= 1


def test_unsafe_archives_limits_and_skips(search_client, monkeypatch):
    client, cid, index = search_client
    root = ROOT / "archives"
    root.mkdir()
    archive(root / "bad.zip", "MARE-001", "needle")
    with zipfile.ZipFile(root / "bad.zip", "a") as z:
        z.writestr("../escape.txt", "needle")
    archive(root / "symlink.zip", "MARE-001", "needle")
    with zipfile.ZipFile(root / "symlink.zip", "a") as z:
        info = zipfile.ZipInfo("pkg/link")
        info.external_attr = (stat.S_IFLNK | 0o777) << 16
        z.writestr(info, "target")
    folder = ROOT / "cases" / str(cid) / "mip/supporting"
    (folder / "binary.exe").write_bytes(b"\x00needle")
    (folder / "scanned.pdf").write_bytes(text_pdf(""))
    (folder / "backups").mkdir()
    (folder / "backups/hidden.txt").write_text("secret needle")
    (folder / "session.key").write_text("secret needle")
    refresh(index)
    with app.app_context():
        run("UPDATE users SET role='Administrator'")
    status = results(client, "needle")["status"]
    reasons = " ".join(x["reason"] for x in status["skipped"])
    assert (
        "unsafe" in reasons
        and "Unsupported format" in reasons
        and "no extractable text" in reasons
    )
    assert results(client, "secret needle")["artifacts"] == 0
    assert extract(b"x" * (MAX_BYTES + 1), "big.txt")[1].startswith("Oversized")
    import mip_search

    monkeypatch.setattr(mip_search, "MAX_MEMBERS", 1)
    refresh(index)
    assert any("4,000 members" in x["reason"] for x in index.status()["skipped"])


def test_auth_navigation_and_permissions(search_client):
    client, cid, index = search_client
    assert app.test_client().get("/search").status_code == 302
    assert app.test_client().get("/search/results?q=x").status_code == 302
    assert app.test_client().get("/search/findings/anything").status_code == 302
    html = client.get("/search").data.decode()
    assert html.index("MIP BOARD") < html.index("SEARCH") < html.index("METRICS")
    refresh(index)
    r = results(client, "case needle")
    row = findings(r)[0]
    with app.app_context():
        run("INSERT INTO roles VALUES('Restricted','',0)")
        run("UPDATE users SET role='Restricted'")
    assert client.get("/search").status_code == 403
    assert client.get("/search/findings/" + row["id"]).status_code == 403
    with app.app_context():
        run("INSERT INTO role_permissions VALUES('Restricted','cases:read')")
    assert client.get("/search").status_code == 200
    assert results(client, "case needle")["artifacts"] == 1


def test_pagination_and_update_after_edit(search_client):
    client, cid, index = search_client
    with app.app_context():
        for i in range(11):
            run(
                "INSERT INTO cases(number,name,stage,created) VALUES(?, 'shared needle',1,'now')",
                (f"CASE-{i}",),
            )
    refresh(index)
    r = results(client, "shared needle")
    assert r["packages"] == 11 and len(r["groups"]) == 10 and r["pages"] == 2
    assert len(results(client, "shared needle", page=2)["groups"]) == 1
    before = index.invalidation.stat().st_mtime_ns if index.invalidation.exists() else 0
    response = client.post(
        f"/cases/{cid}/notes",
        data={"csrf": "search-token", "title": "Added", "body": "new indexed needle"},
    )
    assert response.status_code in (200, 302), response.data
    assert index.invalidation.stat().st_mtime_ns > before
    refresh(index)
    assert results(client, "new indexed needle", kind="notes")["artifacts"] == 1


def test_upload_edit_import_and_case_deletion_invalidation(search_client):
    client, cid, index = search_client
    upload = client.post(
        f"/cases/{cid}/upload",
        data={
            "csrf": "search-token",
            "category": "supporting",
            "file": (io.BytesIO(b"uploaded unique phrase"), "evidence.txt"),
        },
    )
    assert upload.status_code == 302
    refresh(index)
    row = findings(results(client, "uploaded unique phrase"))[0]
    edit = client.post(
        f"/cases/{cid}/file",
        data={
            "csrf": "search-token",
            "path": "supporting/evidence.txt",
            "content": "edited unique phrase",
        },
    )
    assert edit.status_code == 302
    assert client.get("/search/findings/" + row["id"]).status_code == 409
    refresh(index)
    assert results(client, "uploaded unique phrase")["artifacts"] == 0
    assert results(client, "edited unique phrase")["artifacts"] == 1
    with app.app_context():
        out, name = backup_helpers["case_backup"](cid)
        with out:
            backup = out.read()
    imported = client.post(
        "/import-export/import-raw",
        data={"csrf": "search-token", "backup": (io.BytesIO(backup), "source-raw.zip")},
    )
    assert imported.status_code == 302 and imported.location == "/"
    refresh(index)
    assert results(client, "edited unique phrase")["packages"] == 2
    with app.app_context():
        new = (
            db().execute("SELECT id,number FROM cases WHERE id != ?", (cid,)).fetchone()
        )
    response = client.post(
        f'/cases/{new["id"]}/delete',
        data={"csrf": "search-token", "confirm_number": new["number"]},
    )
    assert response.status_code == 302
    # Even a previous snapshot immediately hides deleted cases.
    assert results(client, "edited unique phrase")["packages"] == 1
    refresh(index)
    assert results(client, "edited unique phrase")["packages"] == 1


def test_generated_uploaded_reports_and_job_completion_invalidation(
    search_client, monkeypatch
):
    import report_routes
    from reporting import TOKENS

    client, cid, index = search_client
    template = Document()
    for key in TOKENS:
        template.add_paragraph("{{ " + key + " }}")
    raw = io.BytesIO()
    template.save(raw)
    response = client.post(
        "/account/templates",
        data={
            "csrf": "search-token",
            "template": (io.BytesIO(raw.getvalue()), "search-template.docx"),
        },
    )
    assert response.status_code in (200, 302)

    class InlineThread:
        def __init__(self, target, args, **kwargs):
            self.target, self.args = target, args

        def start(self):
            self.target(*self.args)

    monkeypatch.setattr(report_routes.threading, "Thread", InlineThread)
    sections = {key: "generated searchable statement" for key in TOKENS}
    response = client.post(
        f"/cases/{cid}/generate-report",
        data={
            "csrf": "search-token",
            "generation_mode": "manual",
            "generate_pdf": "0",
            **sections,
        },
    )
    assert response.status_code == 202, response.data
    assert client.get(f"/cases/{cid}/report-job").json["job"]["status"] == "completed"
    refresh(index)
    assert any(
        f["type"] == "docx"
        for f in findings(
            results(client, "generated searchable statement", kind="reports")
        )
    )
    final = Document()
    final.add_paragraph("uploaded final searchable statement")
    raw = io.BytesIO()
    final.save(raw)
    response = client.post(
        f"/cases/{cid}/final-report",
        data={
            "csrf": "search-token",
            "report": (io.BytesIO(raw.getvalue()), "final.docx"),
            "backup": "1",
            "generate_pdf": "0",
        },
    )
    assert response.status_code == 200, response.data
    refresh(index)
    rows = findings(results(client, "uploaded final searchable statement"))
    assert len(rows) == 1 and rows[0]["type"] == "docx"
    assert not any(
        "backups" in f["label"]
        for f in findings(results(client, "generated searchable statement"))
    )


def test_archive_docx_pdf_nested_and_compression_limits(search_client):
    client, cid, index = search_client
    root = ROOT / "archives"
    root.mkdir()
    archive(root / "documents.zip", "MARE-001", "readable evidence")
    word = Document()
    word.add_paragraph("archived Word paragraph")
    word.add_table(1, 1).cell(0, 0).text = "archived Word table"
    buffer = io.BytesIO()
    word.save(buffer)
    with zipfile.ZipFile(
        root / "documents.zip", "a", compression=zipfile.ZIP_DEFLATED
    ) as z:
        z.writestr("pkg/reports/report.docx", buffer.getvalue())
        z.writestr("pkg/reports/report.pdf", text_pdf("archived PDF evidence"))
        z.writestr("pkg/supporting/nested.zip", b"fake nested content")
        z.writestr("pkg/supporting/bomb.txt", b"x" * (2 * 1024 * 1024))
    refresh(index)
    assert (
        results(client, "archived Word paragraph", source="archive")["artifacts"] == 1
    )
    assert results(client, "archived Word table", source="archive")["artifacts"] == 1
    assert results(client, "archived PDF evidence", source="archive")["artifacts"] == 1
    assert results(client, "fake nested content")["artifacts"] == 0
    assert results(client, "nested.zip", kind="filenames")["artifacts"] == 1
    with app.app_context():
        run("UPDATE users SET role='Administrator'")
    reasons = " ".join(
        item["reason"] for item in results(client, "readable")["status"]["skipped"]
    )
    assert "Nested archives" in reasons and "compression ratio" in reasons


def test_database_stale_and_category_permission_revalidation(search_client):
    client, cid, index = search_client
    with app.app_context():
        nid = run(
            "INSERT INTO notes(case_id,title,body) VALUES(?,'Saved','original note phrase')",
            (cid,),
        ).lastrowid
    refresh(index)
    row = findings(results(client, "original note phrase", kind="notes"))[0]
    with app.app_context():
        run("UPDATE notes SET body=? WHERE id=?", ("changed note", nid))
    assert client.get("/search/findings/" + row["id"]).status_code == 409
    with app.app_context():
        run("INSERT INTO roles VALUES('Read metadata','',0)")
        run("INSERT INTO role_permissions VALUES('Read metadata','cases:read')")
        run("UPDATE users SET role='Read metadata'")
    path = ROOT / "cases" / str(cid) / "mip/reports/sensitive.txt"
    path.write_text("protected report phrase")
    refresh(index)
    assert results(client, "protected report phrase")["artifacts"] == 0
    with app.app_context():
        run("UPDATE users SET role='Administrator'")
    report = findings(results(client, "protected report phrase"))[0]
    with app.app_context():
        run("UPDATE users SET role='Read metadata'")
    assert client.get("/search/findings/" + report["id"]).status_code == 403


def test_literals_and_resource_partial_coverage(search_client, monkeypatch):
    from mip_search import snippets

    assert snippets("a+b[c] a+b[c]", "a+b[c]")[0] == 2
    assert snippets("HXXPS[:]//evil[.]example/@x", "https://evil.example/@x")[0] == 1
    client, cid, index = search_client
    path = ROOT / "cases" / str(cid) / "mip/supporting/oversized.txt"
    path.write_text("bounded text phrase")
    import mip_search

    monkeypatch.setattr(mip_search, "MAX_INDEX_TEXT", 1)
    refresh(index)
    r = results(client, "oversized.txt", kind="filenames")
    assert r["artifacts"] == 1 and "budget" in findings(r)[0]["warning"]
    assert r["status"]["partial"]


def test_zip_downloads_are_not_retained(search_client):
    client, cid, index = search_client
    with app.app_context():
        run("UPDATE cases SET stage=2 WHERE id=?", (cid,))
    (ROOT / "cases" / str(cid) / "mip/reports/malware-analysis-report.md").write_text(
        "Completed analysis report"
    )
    before = set(ROOT.rglob("*.zip"))
    response = client.post(
        f"/cases/{cid}/archive",
        data={"csrf": "search-token", "archive_type": "raw", "acknowledge": "1"},
    )
    assert response.status_code == 200, response.data
    assert set(ROOT.rglob("*.zip")) == before
    refresh(index)
    assert results(client, "Search case", source="archive")["artifacts"] == 0


def test_archive_replacement_during_indexing_is_not_published(
    search_client, monkeypatch
):
    import mip_search

    client, cid, index = search_client
    root = ROOT / "archives"
    root.mkdir()
    path = root / "changing.zip"
    archive(path, "MARE-001", "original race content")
    original = mip_search.extract

    def replace_while_extracting(raw, name):
        if name == "pkg/supporting/evidence.txt":
            replacement = root / "replacement.tmp"
            archive(replacement, "MARE-001", "replacement race content")
            replacement.replace(path)
        return original(raw, name)

    monkeypatch.setattr(mip_search, "extract", replace_while_extracting)
    refresh(index)
    assert results(client, "original race content")["artifacts"] == 0
    assert any(
        "changed during indexing" in item["reason"]
        for item in index.status()["skipped"]
    )


def test_refresh_preserves_workflow_database_and_workspace(search_client):
    client, cid, index = search_client
    with app.app_context():
        before = list(db().iterdump())
    files = {
        p.relative_to(ROOT): p.read_bytes()
        for p in (ROOT / "cases").rglob("*")
        if p.is_file()
    }
    refresh(index)
    with app.app_context():
        assert list(db().iterdump()) == before
    assert files == {
        p.relative_to(ROOT): p.read_bytes()
        for p in (ROOT / "cases").rglob("*")
        if p.is_file()
    }


def test_indicator_reference_task_updates_and_csv_mirrors(search_client):
    client, cid, index = search_client

    def post(path, **values):
        return client.post(
            f"/cases/{cid}/{path}", data={"csrf": "search-token", **values}
        )

    assert (
        post(
            "indicators",
            columns=json.dumps(["type", "value", "description"]),
            rows=json.dumps([["domain", "old[.]evil[.]example", "indicator content"]]),
            revision="0",
        ).status_code
        == 200
    )
    assert (
        post(
            "references",
            rows=json.dumps([["https://old.evil.example", "reference content"]]),
            revision="0",
        ).status_code
        == 200
    )
    assert (
        post("tasks", title="Analyst task", body="old task content").status_code == 302
    )
    refresh(index)
    assert results(client, "old.evil.example", kind="indicators")["artifacts"] == 1
    assert results(client, "iocs/indicators.csv", kind="filenames")["artifacts"] == 1
    assert (
        results(client, "supporting/references.csv", kind="filenames")["artifacts"] == 1
    )
    assert results(client, "old.evil.example")["artifacts"] == 2
    with app.app_context():
        task = (
            db().execute("SELECT id FROM tasks WHERE case_id=?", (cid,)).fetchone()[0]
        )
    assert (
        post(
            f"tasks/{task}/edit",
            title="Updated task",
            body="new task content",
            priority="High",
            state="Completed",
        ).status_code
        == 302
    )
    assert (
        post(
            "indicators",
            columns=json.dumps(["type", "value", "description"]),
            rows=json.dumps([["domain", "new[.]evil[.]example", "updated indicator"]]),
            revision="1",
        ).status_code
        == 200
    )
    assert (
        post(
            "references",
            rows=json.dumps([["https://new.evil.example", "updated reference"]]),
            revision="1",
        ).status_code
        == 200
    )
    refresh(index)
    assert results(client, "old.evil.example")["artifacts"] == 0
    assert results(client, "new.evil.example")["artifacts"] == 2
    assert results(client, "new task content", kind="tasks")["artifacts"] == 1
