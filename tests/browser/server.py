"""Isolated browser-test fixture server. Never imported by the production app."""
import os
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
if not os.environ.get('MARE_EDITOR_BROWSER_TEST') or not os.environ.get('MARE_DATA_DIR'):
    raise RuntimeError('Browser tests require an explicitly isolated data directory')
from app import app, run, db, ROOT, ensure_package, package, case
from reporting import TOKENS, current_report
from werkzeug.security import generate_password_hash
from docx import Document

with app.app_context():
    uid = run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",
              ('editor-test', generate_password_hash('temporary-editor-test-123'), 'Administrator')).lastrowid
    cid = run("INSERT INTO cases(number,name,created,stage,owner,creator) VALUES('CM-TEST','Editor regression','2026',1,?,?)", (uid, uid)).lastrowid
    root = ensure_package(cid)
    initial = '# Saved report\n\nExisting **Markdown** and `code`.\n\n::: align-center\nPreserved alignment\n:::\n'
    (root/'reports/sections').mkdir(parents=True, exist_ok=True)
    (root/'reports/sections/executive_summary.md').write_text(initial)
    task_id = run("INSERT INTO tasks(case_id,title,body,created) VALUES(?,?,?,?)", (cid, 'Existing task', 'Saved task **body**', '2026')).lastrowid
    template = Document()
    for key, label in TOKENS.items():
        template.add_heading(label, 1)
        template.add_paragraph('{{ '+key+' }}')
    template_path = ROOT/'editor-test-template.docx'
    template.save(template_path)

# Explicit test-only installed snapshot for management Enable/Disable integration.
if os.environ.get('MARE_PLUGIN_BROWSER_TEST'):
    import plugin_store
    from plugin_discovery import discover
    packages,_=discover(os.environ['MARE_PLUGINS_DIR'])
    with app.app_context():
        owner=db().execute('SELECT user_uuid FROM users WHERE id=?',(uid,)).fetchone()[0]
        for location,entries,manifest in packages:
            timestamp=plugin_store.now();pid=manifest['id']
            db().execute('INSERT INTO developer_plugins(id,owner_uuid,name,created_at,updated_at,published_version,source_location) VALUES(?,?,?,?,?,?,?)',(pid,owner,manifest['name'],timestamp,timestamp,manifest['version'],location))
            db().executemany('INSERT INTO developer_files VALUES(?,?,?,?)',[(pid,path,item['body'],item['directory']) for path,item in entries.items()])
            db().execute('INSERT INTO developer_versions VALUES(?,?,?,?,?,?,?)',(pid,manifest['version'],plugin_store.fingerprint(entries),plugin_store.archive(entries),plugin_store.dumps(manifest),owner,timestamp))
        db().commit()

@app.get('/__editor_test__')
def editor_fixture():
    with app.app_context():
        connection = db()
        report = package(cid)/current_report(case(cid), package(cid))
        document = Document(report) if report.exists() else None
        paragraphs = [p.text for p in document.paragraphs] if document else []
        inline_shapes = len(document.inline_shapes) if document else 0
    return dict(cid=cid, task_id=task_id, initial=initial, template=str(template_path), paragraphs=paragraphs, inline_shapes=inline_shapes)

if __name__ == '__main__':
    if os.environ.get('MARE_EDITOR_TEST_SCRIPT_PREFIX'):
        from werkzeug.middleware.dispatcher import DispatcherMiddleware
        original = app.wsgi_app
        app.wsgi_app = DispatcherMiddleware(original, {os.environ['MARE_EDITOR_TEST_SCRIPT_PREFIX']: original})
    app.run(host='127.0.0.1', port=int(os.environ['MARE_EDITOR_TEST_PORT']), use_reloader=False)
