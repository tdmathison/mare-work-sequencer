"""Server/data compatibility checks; geometry and interaction live in browser tests."""
from html.parser import HTMLParser
from app import app, db, run, ensure_package


class Sources(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.values = {}
        self.current = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'textarea' and 'data-markdown-editor' in attrs:
            self.current = attrs['name']
            self.values[self.current] = ''

    def handle_data(self, value):
        if self.current is not None:
            self.values[self.current] += value

    def handle_endtag(self, tag):
        if tag == 'textarea':
            self.current = None


def fixture():
    with app.app_context():
        uid = run("INSERT INTO users(username,password,role,forced) VALUES('cm-user','unused','User',0)").lastrowid
        cid = run("INSERT INTO cases(number,name,created,stage) VALUES('CM-1','Editors','2026',1)").lastrowid
        root = ensure_package(cid)
        (root/'reports/sections').mkdir(exist_ok=True)
        text = '# Preserved\n\n**bold** `code` <script>literal</script>\n\n::: figure\n![Image](assets/existing.png)\n*Figure 1:*\n:::\n'
        (root/'reports/sections/executive_summary.md').write_text(text)
        run("INSERT INTO tasks(case_id,title,body,created) VALUES(?,?,?,?)", (cid, 'Task', 'Original **task** Markdown', '2026'))
    client = app.test_client()
    with client.session_transaction() as session:
        session.update(uid=uid, version=1, csrf='cm-token')
    return client, cid, root, text


def test_editor_mounts_preserve_markdown_and_database():
    client, cid, root, text = fixture()
    with app.app_context():
        before = '\n'.join(db().iterdump())
    page = client.get(f'/cases/{cid}').data.decode()
    parser = Sources()
    parser.feed(page)
    assert parser.values['executive_summary'].lstrip('\n') == text
    assert parser.values['description'] == 'Original **task** Markdown'
    assert 'vendor/codemirror/markdown-editor.js' in page
    assert 'vendor/codemirror/markdown-editor.css' in page
    for legacy in ('markdown-highlight', 'line-numbers', 'markdown-syntax.js', 'markdown-gutter.js', 'clipboard-images.js'):
        assert legacy not in page
    assert client.get(f'/cases/{cid}/report-data').json['sections']['executive_summary'] == text
    assert (root/'reports/sections/executive_summary.md').read_text() == text
    with app.app_context():
        assert '\n'.join(db().iterdump()) == before


def test_markdown_file_editing_keeps_route_and_plain_storage():
    client, cid, root, text = fixture()
    path = 'reports/sections/executive_summary.md'
    page = client.get(f'/cases/{cid}/file', query_string={'path': path}).data.decode()
    assert 'data-markdown-editor' in page
    assert client.post(f'/cases/{cid}/file', data={'csrf':'cm-token', 'path':path, 'content':text+'More Markdown'}).status_code == 302
    assert (root/path).read_text() == text+'More Markdown'
    code = root/'scripts/test.py'
    code.write_text('print(1)')
    page = client.get(f'/cases/{cid}/file', query_string={'path':'scripts/test.py'}).data.decode()
    assert '<textarea class="editor"  name="content"' in page
    assert 'data-markdown-editor' not in page


def test_locked_case_retains_existing_authorization():
    client, cid, _, text = fixture()
    with app.app_context():
        run('UPDATE cases SET stage=3 WHERE id=?', (cid,))
    page = client.get(f'/cases/{cid}').data
    assert b'data-markdown-editor' in page and b'workflow-controls" disabled' in page
    assert client.post(f'/cases/{cid}/report-sections', data={'csrf':'cm-token', 'executive_summary':text}).status_code == 409
    assert client.post(f'/cases/{cid}/report-sections', data={'executive_summary':text}).status_code == 403


def test_local_compiled_assets_available():
    for filename, mimetype in [('markdown-editor.js','text/javascript'), ('markdown-editor.css','text/css')]:
        response = app.test_client().get('/static/vendor/codemirror/'+filename)
        assert response.status_code == 200 and response.mimetype == mimetype


def test_word_prose_default_and_explicit_template_typography(tmp_path):
    from docx import Document
    from reporting import TOKENS, fill_template
    for explicit in (None, 'Arial'):
        template = Document()
        if explicit:
            template.styles['Normal'].font.name = explicit
        for key in TOKENS:
            template.add_paragraph('{{ '+key+' }}')
        path = tmp_path/'template.docx'
        output = tmp_path/'report.docx'
        template.save(path)
        sections = {key: 'Prose **bold** and `inline code`.\n\n```text\nsource code\n```' for key in TOKENS}
        fill_template(path, output, sections, '1', 'Typography')
        report = Document(output)
        assert report.styles['Normal'].font.name == (explicit or 'Aptos')
        prose = next(p for p in report.paragraphs if p.text.startswith('Prose'))
        assert next(r for r in prose.runs if r.text == 'inline code').font.name == 'Consolas'
        assert any(r.font.name == 'Consolas' for p in report.paragraphs for r in p.runs if 'source code' in r.text)
