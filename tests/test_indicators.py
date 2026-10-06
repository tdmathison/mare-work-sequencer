import io,json,subprocess
from pathlib import Path
from docx import Document
from app import app,db,run,package,ensure_package
from werkzeug.security import generate_password_hash
from reporting import TOKENS,OUTPUT
import indicators,report_routes

def test_indicator_defang_refang_rules():
 source=Path('static/indicators.js').read_text()
 script='''const virtualMachine=require('vm'),assertions=require('assert');
 const context={window:{},document:{getElementById:()=>null}};
 virtualMachine.runInNewContext(SOURCE,context);
 const values=context.window.MareIndicatorValues;
 assertions.equal(values.defang('192.168.1.1'),'192.168.1[.]1');
 assertions.equal(values.defang('evil.example.com'),'evil.example[.]com');
 assertions.equal(values.defang('evil[.]example.com'),'evil[.]example[.]com');
 assertions.equal(values.defang('https://evil.example.com/path'),'hxxps://evil.example[.]com/path');
 assertions.equal(values.defang('http://192.168.1.1:80/path'),'hxxp://192.168.1[.]1:80/path');
 assertions.equal(values.refang('hxxps://evil[.]example[.]com'),'https://evil.example.com');
 assertions.equal(values.refang('192.168[.]1[.]1'),'192.168.1.1');
 assertions.equal(values.refang('ordinary-value'),'ordinary-value');'''.replace('SOURCE',json.dumps(source))
 subprocess.run(['node','-e',script],check=True,capture_output=True,text=True)

def setup():
    with app.app_context():
        uid=run("INSERT INTO users(username,password,role,forced) VALUES(?,?,?,0)",('ioc-user',generate_password_hash('password-long-123'),'User')).lastrowid
        cid=run("INSERT INTO cases(number,name,created,stage) VALUES('IOC-1','IOC case','2026',1)").lastrowid
        ensure_package(cid)
    client=app.test_client()
    with client.session_transaction() as session:session.update(uid=uid,version=1,csrf='ioc-token')
    def post(path,**data):return client.post(path,data={'csrf':'ioc-token',**data})
    return client,cid,post

def test_fixed_table_csv_append_and_locks():
    client,cid,post=setup();base=f'/cases/{cid}/indicators'
    columns=['type','value','description']
    assert client.get(base).json['columns']==columns
    saved=post(base,rows=json.dumps([['domain','evil.example.com','C2']]),revision='0');assert saved.status_code==200
    # Headerless pasted rows use the fixed order, with quoted commas intact.
    assert post(base+'/import',csv='IP,192.0.2.1,"C2, observed"\n',revision='1').json['added']==1
    # File columns may be reordered, mixed case, and include ignored extra fields.
    imported=post(base+'/import',csv=(io.BytesIO(b'Extra,DESCRIPTION, VaLuE ,TYPE\nIgnored,Payload,'+b'a'*64+b',SHA256\n'),'iocs.csv'),revision='2')
    assert imported.status_code==200 and len(imported.json['rows'])==3
    assert imported.json['rows'][1][2]=='C2, observed'
    assert imported.json['rows'][2]==['SHA256','a'*64,'Payload']
    pasted=post(base+'/import',csv='Description,Type,Value\nDNS,domain,next.example.com\n',revision='3')
    assert pasted.status_code==200 and pasted.json['rows'][-1]==['domain','next.example.com','DNS']
    for text in ['Value,Type\nx,y','type,value,description\nx,y','type,value,value,description\nx,y,z,q']:
        assert post(base+'/import',csv=text,revision='4').status_code==400
    assert post(base+'/import',csv=(io.BytesIO(b'IP,192.0.2.1,Observed'),'missing-header.csv'),revision='4').status_code==400
    assert post(base,rows='[]',revision='0').status_code==409
    assert post(base,columns=json.dumps(['kind','value','description']),rows='[]',revision='4').status_code==409
    assert len(client.get(base).json['rows'])==4
    result=post(base,rows=json.dumps(pasted.json['rows'][1:]),revision='4')
    assert result.status_code==200 and len(result.json['rows'])==3
    assert '192[.]0[.]2[.]1' in (package(cid)/'iocs/indicators.csv').read_text()
    page=client.get(f'/cases/{cid}').data
    assert b'tab-indicators' in page and b'add-indicator-column' not in page
    assert b'name="indicators_of_compromise"' not in page and b'name="mitre_attack_mapping"' not in page
    with app.app_context():run('UPDATE cases SET stage=3 WHERE id=?',(cid,))
    assert post(base,rows='[]',revision='5').status_code==409
    assert post(base+'/import',csv='type,value,description\n',revision='5').status_code==409

def test_word_table_uses_fixed_columns_and_plain_text(monkeypatch,tmp_path):
    client,cid,post=setup()
    columns=['type','value','description'];rows=[['domain','evil.example.com','literal **text** | evidence']]
    post(f'/cases/{cid}/indicators',columns=json.dumps(columns),rows=json.dumps(rows),revision='0')
    doc=Document()
    for key,label in TOKENS.items():doc.add_heading(label,1);doc.add_paragraph('{{ '+key+' }}')
    template=tmp_path/'template.docx';doc.save(template)
    post('/account/templates',template=(io.BytesIO(template.read_bytes()),'template.docx'))
    class InlineThread:
        def __init__(self,target,args,**kwargs):self.target=target;self.args=args
        def start(self):self.target(*self.args)
    monkeypatch.setattr(report_routes.threading,'Thread',InlineThread)
    result=post(f'/cases/{cid}/generate-report',generation_mode='manual',executive_summary='Summary')
    assert result.status_code==202
    assert client.get(f'/cases/{cid}/report-job').json['job']['status']=='completed'
    doc=Document(package(cid)/client.get(f'/cases/{cid}/report-job').json['report_path']);table=next(t for t in doc.tables if [c.text for c in t.rows[0].cells]==columns)
    assert [c.text for c in table.rows[1].cells]==['domain','evil[.]example[.]com','literal **text** | evidence']

def test_existing_custom_table_migration_preserves_original():
    client,cid,post=setup()
    with app.app_context():
        connection=db()
        columns=['Type','Value','Description','Source'];rows=[['domain','old.example.com','C2','Sandbox']]
        run('INSERT INTO indicator_tables VALUES(?,?,?,?)',(cid,json.dumps(columns),json.dumps(rows),2))
        assert indicators.migrate_fixed_columns(connection)==[cid]
        table=indicators.load_table(db,cid)
        assert table['columns']==['type','value','description'] and table['rows']==[rows[0][:3]]
        original=connection.execute('SELECT * FROM indicator_table_legacy WHERE case_id=?',(cid,)).fetchone()
        assert json.loads(original['rows'])==rows and json.loads(original['columns'])==columns
        assert indicators.migrate_fixed_columns(connection)==[]

def test_indicator_sections_save_import_revision_and_existing_rows():
    import csv
    client,cid,post=setup();base=f'/cases/{cid}/indicators'
    post(base,rows=json.dumps([['domain','old.example.com','Previous entry']]),revision='0')
    initial=client.get(base).json
    assert initial['sections']==[{'id':'default','title':indicators.DEFAULT_SECTION_TITLE,'description':indicators.DEFAULT_SECTION_DESCRIPTION,'rows':initial['rows']}]
    sections=initial['sections']+[{'id':'payload','title':'Payload hashes','description':'Files recovered from the sandbox.','rows':[['sha256','a'*64,'Payload']]}]
    saved=post(base,sections=json.dumps(sections),revision='1');assert saved.status_code==200
    assert saved.json['sections']==sections and len(saved.json['rows'])==2
    assert post(base,sections=json.dumps(sections),revision='1').status_code==409
    assert post(base,rows='[]',revision='2').status_code==409 # Old clients cannot flatten away grouping.
    imported=post(base+'/import',section_id='payload',csv=(io.BytesIO(b'Description,VALUE,Type\nC2,192.0.2.1,ipv4\n'),'iocs.csv'),revision='2')
    assert imported.status_code==200 and imported.json['added']==1
    assert len(imported.json['sections'][0]['rows'])==1 and len(imported.json['sections'][1]['rows'])==2
    assert imported.json['sections'][1]['rows'][-1]==['ipv4','192.0.2.1','C2']
    assert post(base+'/import',section_id='missing',csv='domain,test.example,Test',revision='3').status_code==400
    assert post(base+'/import',section_id='default',csv='domain,test.example,Test',revision='2').status_code==400
    exported=list(csv.DictReader(io.StringIO((package(cid)/'iocs/indicators.csv').read_text())))
    assert exported[0]['section']==indicators.DEFAULT_SECTION_TITLE and exported[1]['section']=='Payload hashes'
    assert exported[2]['value']=='192[.]0[.]2[.]1' and exported[2]['section_description']=='Files recovered from the sandbox.'
    invalid=[{**sections[0],'title':''},sections[1]]
    assert post(base,sections=json.dumps(invalid),revision='3').status_code==409
    assert client.get(base).json['revision']==3
    page=client.get(f'/cases/{cid}').data.decode()
    assert page.index('id="destructive-commands"')<page.index('id="panel-tasks"')
    assert 'id="add-indicator-section"' in page
    with app.app_context():run('UPDATE cases SET stage=3 WHERE id=?',(cid,))
    assert post(base,sections=json.dumps(sections),revision='3').status_code==409


def test_word_report_indicator_subsections(monkeypatch,tmp_path):
    client,cid,post=setup();base=f'/cases/{cid}'
    sections=[{'id':'default','title':'Default','description':'Initial observations.','rows':[['domain','evil.example.com','literal **C2**']]},{'id':'payload','title':'Payload hashes','description':'Recovered files | sandbox evidence.','rows':[['sha256','b'*64,'Payload']]},{'id':'empty','title':'Other infrastructure','description':'No additional infrastructure identified.','rows':[]}]
    assert post(base+'/indicators',sections=json.dumps(sections),revision='0').status_code==200
    doc=Document()
    for key,label in TOKENS.items():doc.add_heading(label,1);doc.add_paragraph('{{ '+key+' }}')
    template=tmp_path/'sections.docx';doc.save(template)
    post('/account/templates',template=(io.BytesIO(template.read_bytes()),'sections.docx'))
    class InlineThread:
        def __init__(self,target,args,**kwargs):self.target=target;self.args=args
        def start(self):self.target(*self.args)
    monkeypatch.setattr(report_routes.threading,'Thread',InlineThread)
    assert post(base+'/generate-report',executive_summary='Summary').status_code==202
    assert client.get(base+'/report-job').json['job']['status']=='completed'
    doc=Document(package(cid)/client.get(f'/cases/{cid}/report-job').json['report_path'])
    heading_text=[p.text for p in doc.paragraphs if p.style.name=='Heading 2']
    assert heading_text[:3]==['Default','Payload hashes','Other infrastructure']
    assert all(s['description'] in [p.text for p in doc.paragraphs] for s in sections)
    tables=[t for t in doc.tables if [c.text for c in t.rows[0].cells]==indicators.DEFAULT_COLUMNS]
    assert len(tables)==3
    assert [c.text for c in tables[0].rows[1].cells]==['domain','evil[.]example[.]com','literal **C2**']
    assert len(tables[2].rows)==1
    with app.app_context():number=db().execute('SELECT number FROM cases WHERE id=?',(cid,)).fetchone()[0]
    assert post(base+'/delete',confirm_number=number).status_code==302
    with app.app_context():assert not db().execute('SELECT 1 FROM indicator_sections WHERE case_id=?',(cid,)).fetchone()

def test_section_edit_delete_and_protected_default():
    client,cid,post=setup();base=f'/cases/{cid}/indicators'
    initial=client.get(base).json
    assert initial['sections'][0]['title']==indicators.DEFAULT_SECTION_TITLE
    assert initial['sections'][0]['description']==indicators.DEFAULT_SECTION_DESCRIPTION
    sections=initial['sections']+[{'id':'additional','title':'Infrastructure','description':'C2 servers.','rows':[['ipv4','192.0.2.1','C2']]}]
    assert post(base,sections=json.dumps(sections),revision='0').status_code==200
    sections[0]['title']='Analyst-selected indicators';sections[0]['description']='Confirmed findings.'
    sections[1]['title']='Updated infrastructure';sections[1]['description']='Revised findings.'
    result=post(base,sections=json.dumps(sections),revision='1')
    assert result.status_code==200 and result.json['sections']==sections
    rejected=post(base,sections=json.dumps(sections[1:]),revision='2')
    assert rejected.status_code==409 and 'cannot be deleted' in rejected.json['error']
    assert client.get(base).json['sections']==sections
    saved=post(base,sections=json.dumps(sections[:1]),revision='2')
    assert saved.status_code==200 and saved.json['sections']==sections[:1]
    assert saved.json['rows']==[]
    assert '192[.]0[.]2[.]1' not in (package(cid)/'iocs/indicators.csv').read_text()
    with app.app_context():
        assert indicators.migrate_section_defaults(db())==[] # Preserve analyst overrides on restart.


def test_section_defaults_migration_preserves_rows_and_overrides():
    client,cid,post=setup();base=f'/cases/{cid}/indicators'
    sections=[{'id':'default','title':'Default','description':'','rows':[['domain','old.example.com','Existing']]}]
    assert post(base,sections=json.dumps(sections),revision='0').status_code==200
    with app.app_context():
        assert indicators.migrate_section_defaults(db())==[cid]
        current=indicators.load_table(db,cid)
        assert current['sections'][0]['title']==indicators.DEFAULT_SECTION_TITLE
        assert current['sections'][0]['description']==indicators.DEFAULT_SECTION_DESCRIPTION
        assert current['rows']==sections[0]['rows'] and current['sections'][0]['rows']==sections[0]['rows']
        assert current['revision']==2
        assert indicators.migrate_section_defaults(db())==[]
    current['sections'][0]['title']='Customized analysis indicators';current['sections'][0]['description']='Analyst context.'
    assert post(base,sections=json.dumps(current['sections']),revision='2').status_code==200
    with app.app_context():
        assert indicators.migrate_section_defaults(db())==[]
        assert indicators.load_table(db,cid)['sections'][0]['description']=='Analyst context.'

    # Later deliberate overrides matching the old stock wording also survive restart.
    current['sections'][0]['title']='Default';current['sections'][0]['description']=''
    assert post(base,sections=json.dumps(current['sections']),revision='3').status_code==200
    with app.app_context():
        assert indicators.migrate_section_defaults(db())==[]
        assert indicators.load_table(db,cid)['sections'][0]['title']=='Default'
