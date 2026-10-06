"""Report extraction and Word rendering. Uploaded content is data, never executed."""
import os, re, shutil, subprocess, tempfile, zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SECTIONS={'executive_summary':'Executive Summary','key_findings':'Key Findings','detection_opportunities':'Detection Opportunities','reverse_engineering_findings':'Reverse Engineering Findings'}
GENERATED={'mitre_attack_mapping':'MITRE ATT&CK Mapping','mitre_mbc_mapping':'MITRE MBC Mapping','indicators_of_compromise':'Indicators of Compromise','appendices':'Appendices'}
TOKENS={**SECTIONS,**GENERATED}
OUTPUT='reports/malware-analysis-report.docx'

def convert_docx_to_pdf(source,destination):
    executable=shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:raise ValueError('PDF conversion unavailable: LibreOffice Writer was not found on PATH (expected libreoffice or soffice).')
    source=Path(source);destination=Path(destination)
    destination.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='mare-pdf-',dir=destination.parent) as directory:
        temporary=Path(directory);output=temporary/'output';output.mkdir()
        command=[executable,'-env:UserInstallation='+(temporary/'profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(output),str(source)]
        try:result=subprocess.run(command,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=120)
        except (OSError,subprocess.TimeoutExpired) as exc:raise ValueError('LibreOffice PDF conversion failed: '+str(exc)[:300]) from exc
        converted=output/(source.stem+'.pdf')
        if result.returncode or not converted.is_file() or not converted.stat().st_size:
            detail=(result.stderr or result.stdout).strip()
            raise ValueError('LibreOffice PDF conversion failed'+(': '+detail[:300] if detail else ' (exit '+str(result.returncode)+', no output PDF).'))
        os.replace(converted,destination)

def report_filename(c,date=None):
    from datetime import datetime,timezone
    def component(value,limit):
        value=re.sub(r'[^\w-]+','_',str(value),flags=re.UNICODE)
        value=re.sub(r'_+','_',value).strip('_')
        # Bound UTF-8 bytes for portable filesystem component lengths.
        return value.encode('utf-8')[:limit].decode('utf-8',errors='ignore').rstrip('_')
    number=(c['external_number'] or '').strip()
    if not number:
        number=c['number']
        match=re.fullmatch(r'MARE-\d{4}-(\d+)',number)
        if match:number=match.group(1)
    title=re.sub(r'^\d{8}:\s*','',c['name'].strip())
    stamp=date or datetime.now(timezone.utc).strftime('%Y%m%d')
    return stamp+'-MARE_'+(component(number,64) or 'case')+'_RE_Report_'+(component(title,144) or 'case')+'.docx'

def valid_report_path(value):
    return isinstance(value,str) and value.startswith('reports/') and '/' not in value[8:] and '\\' not in value and value.endswith('.docx') and '..' not in value and len(value.encode('utf-8'))<=263

def current_report(c,root):
    value=c['report_path'] if 'report_path' in c.keys() else ''
    return value if value and valid_report_path(value) else OUTPUT

def is_backup(rel): return any(part.lower() in ('backup','backups','.backups') for part in Path(rel).parts)

def validate_docx(path):
    with zipfile.ZipFile(path) as z:
        if '[Content_Types].xml' not in z.namelist() or 'word/document.xml' not in z.namelist(): raise ValueError('Upload a valid DOCX Word document.')
        if sum(i.file_size for i in z.infolist())>50*1024*1024: raise ValueError('Expanded template is too large.')
        if any('vbaproject' in n.lower() for n in z.namelist()): raise ValueError('Macro-enabled documents are not supported.')

def paragraphs(doc):
    for p in doc.paragraphs: yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs: yield p

def template_sections(path):
    validate_docx(path);doc=Document(path);found=set()
    for p in paragraphs(doc):
        for key,label in TOKENS.items():
            if key=='key_findings' and (p.text.strip().lower()=='threat overview' or re.fullmatch(r'\s*\{\{\s*threat_overview\s*\}\}\s*',p.text)):found.add(key)
            elif re.fullmatch(r'\s*\{\{\s*'+key+r'\s*\}\}\s*',p.text): found.add(key)
            elif p.text.strip().lower() in (label.lower(),label.lower().replace('att&ck','attack'),label.lower().replace('mbc','mbac')): found.add(key)
    return found

def refang(value): return value.replace('[.]','.').replace('hxxps://','https://').replace('hxxp://','http://').replace('[@]','@').replace('[:] ',':').replace('[:]',':')
def defang(value,kind=''):
    value=refang(str(value)).strip()
    if re.fullmatch(r'[A-Fa-f0-9]{32}|[A-Fa-f0-9]{40}|[A-Fa-f0-9]{64}',value): return value
    value=re.sub(r'(?i)https://','hxxps://',value);value=re.sub(r'(?i)http://','hxxp://',value)
    return value.replace('.','[.]').replace('@','[@]').replace(':','[:]')

def appendix_links(sections):
    links=[];seen=set()
    for text in sections.values():
        for url in re.findall(r'https?://[^\s<>"\]]+',text):
            url=url.rstrip(').,;')
            if url not in seen: links.append(url);seen.add(url)
    return links

def defang_in_text(text):
    return re.sub(r'(?:https?://|hxxps?://)[^\s<>]+|\b(?:[A-Za-z0-9_-]+\.)+[A-Za-z]{2,63}\b|\b(?:\d{1,3}\.){3}\d{1,3}\b',lambda m:defang(m.group()),text)

def md_table(headers,rows):
    def cell(v):return str(v).replace('|','\\|').replace('\n',' ')
    return '| '+' | '.join(headers)+' |\n| '+' | '.join('---' for _ in headers)+' |\n'+''.join('| '+' | '.join(cell(v) for v in row)+' |\n' for row in rows)

def style_report_table(doc,table):
    """Match the supplied Grid Table 4 sample without changing document branding."""
    from docx.oxml import parse_xml
    from docx.shared import Pt,RGBColor
    if not any(node.get(qn('w:styleId'))=='MAREReportTable' for node in doc.styles.element):
        doc.styles.element.append(parse_xml((Path(__file__).parent/'resources/report-table-style.xml').read_bytes()))
    # Aptos font metadata supplies a sans-serif fallback on systems without Aptos.
    try:font_part=doc.part.part_related_by(RT.FONT_TABLE)
    except KeyError:
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        font_part=Part(PackURI('/word/fontTable.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml',b'<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>',doc.part.package)
        doc.part.relate_to(font_part,RT.FONT_TABLE)
    fonts_root=parse_xml(font_part.blob)
    if not any(node.get(qn('w:name'))=='Aptos' for node in fonts_root):
        from lxml import etree
        fonts_root.append(parse_xml((Path(__file__).parent/'resources/report-table-font.xml').read_bytes()))
        font_part._blob=etree.tostring(fonts_root,xml_declaration=True,encoding='UTF-8',standalone=True)
    table.style=doc.styles['MARE Report Table']
    look=table._tbl.tblPr.find(qn('w:tblLook'))
    if look is None:look=OxmlElement('w:tblLook');table._tbl.tblPr.append(look)
    for name,value in {'val':'04A0','firstRow':'1','lastRow':'0','firstColumn':'1','lastColumn':'0','noHBand':'0','noVBand':'1'}.items():look.set(qn('w:'+name),value)
    for ri,row in enumerate(table.rows):
        if ri==0 and row._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                ppr=paragraph._p.get_or_add_pPr()
                rpr=ppr.find(qn('w:rPr'))
                if rpr is None:rpr=OxmlElement('w:rPr');ppr.append(rpr)
                fonts=OxmlElement('w:rFonts')
                for name in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+name),'Aptos')
                rpr.append(fonts)
                paragraph.paragraph_format.space_after=Pt(0)
                paragraph.paragraph_format.line_spacing=1
                for run in paragraph.runs:
                    run.font.name='Aptos'
                    run.font.size=Pt(11)
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor.from_string('FFFFFF')
                    elif ci==0:run.bold=True

# Render Markdown blocks directly into the uploaded document, retaining its styles and relationships.
def inline(paragraph,tokens,assets_root=None):
    bold=False;italic=False;strike=False;link=None
    for t in tokens or []:
        if t.type=='strong_open':bold=True
        elif t.type=='strong_close':bold=False
        elif t.type=='em_open':italic=True
        elif t.type=='em_close':italic=False
        elif t.type=='s_open':strike=True
        elif t.type=='s_close':strike=False
        elif t.type=='html_inline' and '<input' in t.content:
            paragraph.add_run('☑ ' if 'checked' in t.content else '☐ ')
        elif t.type=='link_open':link=t.attrGet('href')
        elif t.type=='link_close':link=None
        elif t.type=='image':
            source=t.attrGet('src') or ''
            if assets_root is not None and re.fullmatch(r'assets/report-image-[a-f0-9]{32}\.png',source):
                path=Path(assets_root)/source
                if not path.is_file():raise ValueError('A report image is missing. Remove its Markdown reference or upload it again.')
                from docx.shared import Inches
                picture=paragraph.add_run().add_picture(str(path),width=Inches(5.5))
                line=OxmlElement('a:ln');line.set('w','9525')
                fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val','444444');fill.append(color);line.append(fill)
                picture._inline.graphic.graphicData.pic.spPr.append(line)
            else:paragraph.add_run(t.content)
        elif t.type in ('text','code_inline','softbreak','hardbreak'):
            text='\n' if t.type in ('softbreak','hardbreak') else t.content
            if link and link.startswith(('http://','https://')):
                h=OxmlElement('w:hyperlink');h.set(qn('r:id'),paragraph.part.relate_to(link,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');props=OxmlElement('w:rPr');style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');props.append(style);r.append(props);node=OxmlElement('w:t');node.text=text;r.append(node);h.append(r);paragraph._p.append(h)
            else:
                r=paragraph.add_run(text);r.bold=bold;r.italic=italic;r.font.strike=strike
                if t.type=='code_inline':r.font.name='Consolas'

def add_markdown(doc,text,assets_root=None):
    tokens=__import__('markdown_support').markdown_parser().parse(text);elements=[];i=0;list_style=None;alignment=[]
    while i<len(tokens):
        t=tokens[i]
        if t.type.startswith('container_figure_'):
            if t.nesting==1:alignment.append('center')
            elif alignment:alignment.pop()
        elif t.type.startswith('container_align-'):
            if t.nesting==1:alignment.append(t.type.split('-')[1].split('_')[0])
            elif alignment:alignment.pop()
        elif t.type in ('bullet_list_open','ordered_list_open'):list_style='List Bullet' if t.type=='bullet_list_open' else 'List Number'
        elif t.type in ('bullet_list_close','ordered_list_close'):list_style=None
        elif t.type in ('paragraph_open','heading_open'):
            style=('Heading '+t.tag[1:]) if t.type=='heading_open' else list_style or 'Normal'
            try:p=doc.add_paragraph(style=style)
            except KeyError:p=doc.add_paragraph()
            if i+1<len(tokens) and tokens[i+1].type=='inline':inline(p,tokens[i+1].children,assets_root)
            if alignment:
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                p.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}[alignment[-1]]
            elements.append(p._p)
        elif t.type in ('fence','code_block'):
            p=doc.add_paragraph();p.add_run(t.content.rstrip()).font.name='Consolas';elements.append(p._p)
        elif t.type=='table_open':
            rows=[];row=[];i+=1
            while i<len(tokens) and tokens[i].type!='table_close':
                if tokens[i].type=='tr_open':row=[]
                elif tokens[i].type=='inline':row.append(tokens[i])
                elif tokens[i].type=='tr_close':rows.append(row)
                i+=1
            if rows:
                table=doc.add_table(rows=len(rows),cols=max(len(r) for r in rows));elements.append(table._tbl)
                for ri,row in enumerate(rows):
                    for ci,token in enumerate(row):inline(table.cell(ri,ci).paragraphs[0],token.children,assets_root)
                style_report_table(doc,table)
        i+=1
    if not elements:elements=[doc.add_paragraph('No applicable findings recorded.')._p]
    return elements

def fill_template(template,destination,sections,case_number,case_name,indicators=None,references=None,assets_root=None,external_system='',external_number=''):
    doc=Document(template)
    anchors={}
    for p in list(paragraphs(doc)):
        for key,label in TOKENS.items():
            if key=='key_findings' and re.fullmatch(r'\s*\{\{\s*threat_overview\s*\}\}\s*',p.text):anchors[key]=(p,True)
            elif key=='key_findings' and p.text.strip().lower()=='threat overview':
                p.text='Key Findings';anchors.setdefault(key,(p,False))
            elif re.fullmatch(r'\s*\{\{\s*'+key+r'\s*\}\}\s*',p.text):anchors[key]=(p,True)
            elif p.text.strip().lower() in (label.lower(),label.lower().replace('att&ck','attack'),label.lower().replace('mbc','mbac')):anchors.setdefault(key,(p,False))
    missing=set(TOKENS)-set(anchors)
    if missing:raise ValueError('Template is missing sections: '+', '.join(sorted(missing)))
    anchor_nodes={p._p for p,_ in anchors.values()}
    for key,(p,marker) in anchors.items():
        node=p._p
        if not marker:
            # Clear the old section body, stopping at the next heading, marker, or section boundary.
            nxt=node.getnext()
            while nxt is not None and nxt.tag!=qn('w:sectPr') and nxt not in anchor_nodes:
                if nxt.tag==qn('w:p'):
                    style=nxt.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
                    if style is not None and ('heading' in str(style.get(qn('w:val'))).lower() or style.get(qn('w:val'))=='Title'):break
                following=nxt.getnext();nxt.getparent().remove(nxt);nxt=following
        if key=='indicators_of_compromise' and indicators is not None and ('sections' in indicators or indicators['rows']):
            blocks=[]
            groups=indicators.get('sections',[{'title':'Default','description':'','rows':indicators['rows']}])
            for section in groups:
                blocks.append(doc.add_heading(section['title'],level=2)._p)
                if section['description'].strip():blocks.append(doc.add_paragraph(section['description'])._p)
                table=doc.add_table(rows=1,cols=len(indicators['columns']))
                for cell,value in zip(table.rows[0].cells,indicators['columns']):cell.text=value
                for row in section['rows']:
                    for cell,value in zip(table.add_row().cells,row):cell.text=value
                style_report_table(doc,table)
                blocks.append(table._tbl)
                if not section['rows']:blocks.append(doc.add_paragraph('No indicators recorded in this section.')._p)
        else: blocks=add_markdown(doc,sections[key],assets_root)
        if key=='appendices' and references is not None:
            heading=doc.add_heading('Links',level=2);blocks.append(heading._p)
            rows=[row for row in references['rows'] if any(v.strip() for v in row)]
            if rows:
                links=doc.add_table(rows=1,cols=2)
                for cell,value in zip(links.rows[0].cells,['Link','Description']):cell.text=value
                for link,description in rows:
                    cells=links.add_row().cells;cells[0].text=link;cells[1].text=description
                    if link.strip().startswith(('https://','http://')):
                        paragraph=cells[0].paragraphs[0];paragraph.clear()
                        hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),paragraph.part.relate_to(link.strip(),RT.HYPERLINK,is_external=True))
                        run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=link;run.append(text);hyperlink.append(run);paragraph._p.append(hyperlink)
                style_report_table(doc,links)
                blocks.append(links._tbl)
            else:blocks.append(doc.add_paragraph('No reference links have been recorded.')._p)
        cursor=node
        if marker:
            for block in blocks:node.addprevious(block)
            node.getparent().remove(node)
        else:
            for block in blocks:cursor.addnext(block);cursor=block
    case_title=re.sub(r'^\d{8}:\s*','',case_name.strip())
    if external_system and external_number:
        display_case_number=f'{external_system} #{external_number}'
    else:
        match=re.fullmatch(r'MARE-(\d{4}-\d+)',case_number)
        display_case_number=f'MARE #{match.group(1)}' if match else case_number
    all_paragraphs=list(paragraphs(doc))
    for section in doc.sections:
        for story in (section.header,section.footer,section.first_page_header,section.first_page_footer,section.even_page_header,section.even_page_footer):
            all_paragraphs.extend(paragraphs(story))
    for p in all_paragraphs:
        for key,value in [('case_number',display_case_number),('case_name',case_name),('case_title',case_title)]:
            if '{{'+key+'}}' in p.text or '{{ '+key+' }}' in p.text:
                p.text=p.text.replace('{{'+key+'}}',value).replace('{{ '+key+' }}',value)
    doc.save(destination)

